"""Módulo Electromecánica: guardado/lectura, validaciones, bloques del informe
y presencia en la presentación semanal. Llama a las funciones de los
endpoints directamente (Mongo local, sin HTTP ni auth real)."""
import asyncio
import io
import os
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_electromecanica")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi import HTTPException  # noqa: E402
from pptx import Presentation  # noqa: E402

import server  # noqa: E402


def _body(**kw):
    base = dict(
        avance={"estaciones": {"fabricacion": 72.1, "suministro": 37.5, "montaje": 0},
                "pilonas": {"fabricacion": 18.94, "suministro": None, "montaje": 0}},
        pilonas_total=60, pilonas_suministradas=10, patio_nota="10 pilonas en patio Cantera",
        documentos=[server.ElectromecDocumento(titulo="Planos de anclaje", estatus="en_revision", nota="Rev. B"),
                    server.ElectromecDocumento(titulo="  ", estatus="recibido")],
        observaciones=[server.ElectromecObservacion(titulo="Galvanizado", texto="Reparar recubrimiento")],
    )
    base.update(kw)
    return server.ElectromecIn(**base)


async def _scenario():
    pid, uid = str(uuid.uuid4()), str(uuid.uuid4())
    user = {"id": uid, "name": "Coord", "role": server.ROLE_COORD, "project_ids": [pid]}
    await server.db.projects.insert_one({
        "id": pid, "name": "Proyecto X", "color_institucional": "#8A1538",
        "modulos": {"reportes_semanales": {"electromecanica": True}},
    })
    try:
        vacio = await server.get_electromecanica(pid, user)
        assert vacio["documentos"] == [] and vacio["avance_general_calculado"] is None

        # validaciones
        for malo in (_body(avance={"estaciones": {"montaje": 120}}),
                     _body(avance={"inventado": {"montaje": 5}}),
                     _body(pilonas_total=5, pilonas_suministradas=9)):
            try:
                await server.set_electromecanica(pid, malo, user)
                raise AssertionError("debió rechazar")
            except HTTPException as e:
                assert e.status_code == 400

        out = await server.set_electromecanica(pid, _body(), user)
        assert len(out["documentos"]) == 1           # el vacío se descarta
        assert out["documentos"][0]["id"]            # ids generados
        # promedio de lo capturado: (72.1+37.5+0+18.94+0)/5
        assert out["avance_general_calculado"] == round((72.1 + 37.5 + 0 + 18.94 + 0) / 5, 2)
        manual = await server.set_electromecanica(pid, _body(avance_general=91.22), user)
        assert manual["avance_general_calculado"] == 91.22

        doc = await server.db.electromecanica.find_one({"project_id": pid})
        bloques = await server._electromec_bloques(doc)
        tipos = [b["type"] for b in bloques]
        assert tipos.count("table") == 1 and tipos.count("imagen") == 3 and tipos[-1] == "list"
        assert bloques[1]["imagen"].startswith("data:image/")

        # Módulo apagado → no sale en el export
        data = await server._gather_modulos_export_data(
            pid, await server.db.projects.find_one({"id": pid}),
            datetime(2020, 1, 1, tzinfo=timezone.utc), datetime(2100, 1, 1, tzinfo=timezone.utc), "reportes_semanales")
        assert data["electromecanica"]["titulo"] == "SISTEMA ELECTROMECÁNICO"
        apagado = dict(await server.db.projects.find_one({"id": pid}))
        apagado["modulos"] = {"reportes_semanales": {}}
        data_off = await server._gather_modulos_export_data(
            pid, apagado, datetime(2020, 1, 1, tzinfo=timezone.utc), datetime(2100, 1, 1, tzinfo=timezone.utc), "reportes_semanales")
        assert data_off["electromecanica"] is None

        # Presentación semanal: divisor + láminas
        fecha = datetime.now(timezone.utc).date().isoformat()
        pres = await server.create_presentacion(
            pid, server.PresentacionIn(project_id=pid, fecha=fecha, plantilla="semanal"), user)
        assert "electromecanico" in [s["tipo"] for s in pres["secciones"]]
        resp = await server.export_presentacion_pptx(pid, pres["id"], user)
        raw = b"".join([c async for c in resp.body_iterator])
        prs = Presentation(io.BytesIO(raw))
        textos = []
        for sl in prs.slides:
            partes = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    partes += [c.text for r in sh.table.rows for c in r.cells]
            textos.append(" ".join(partes))
        todo = " | ".join(textos)
        assert "SISTEMA ELECTROMECÁNICO" in todo and "Planos de anclaje" in todo and "Reparar recubrimiento" in todo
    finally:
        await server.db.projects.delete_one({"id": pid})
        await server.db.electromecanica.delete_many({"project_id": pid})
        await server.db.presentaciones.delete_many({"project_id": pid})


def test_electromecanica_end_to_end():
    asyncio.run(_scenario())
