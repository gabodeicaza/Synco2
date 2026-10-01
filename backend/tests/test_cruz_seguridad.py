"""Cruz de Seguridad (server.py): tablero mensual de seguridad de obra.

Cada día del mes es una casilla de una cruz, pintada según lo peor que pasó:
verde = sin accidentes, gris = día que aún no ocurre. Resumen por tipo,
fuerza de trabajo promedio y bitácora del año — exportables a Excel.
"""
import asyncio
import io
import os
import sys
from datetime import date, datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_cruz")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402

import server  # noqa: E402
from server import _cruz_calcular, _cruz_layout, _cruz_xlsx_bytes  # noqa: E402


def _r(fecha, tipo, **kw):
    return {"fecha": fecha, "tipo_registro": tipo, "created_at": datetime(2026, 1, 1), **kw}


def test_layout_es_una_cruz_de_9_columnas():
    lay = {d: (f, c) for d, f, c in _cruz_layout(30)}
    assert len(lay) == 30
    # brazo superior 3×2, cuerpo 9×2, brazo inferior 3×2 — centrados en las columnas 3-5
    assert [lay[d] for d in (1, 2, 3)] == [(0, 3), (0, 4), (0, 5)]
    assert [lay[d] for d in (4, 5, 6)] == [(1, 3), (1, 4), (1, 5)]
    assert lay[7] == (2, 0) and lay[15] == (2, 8) and lay[16] == (3, 0) and lay[24] == (3, 8)
    assert [lay[d] for d in (25, 26, 27)] == [(4, 3), (4, 4), (4, 5)]
    assert lay[30] == (5, 5)
    # El 31 va debajo del 28; en meses cortos no hay casillas de más.
    assert {d: p for d, *p in [(d, f, c) for d, f, c in _cruz_layout(31)]}[31] == [6, 3]
    assert len(_cruz_layout(28)) == 28 and max(d for d, _, _ in _cruz_layout(29)) == 29


def test_dias_se_pintan_por_el_peor_evento_y_el_futuro_queda_gris():
    regs = [
        _r("2026-03-02", "accidente_leve"),
        _r("2026-03-02", "dano_propiedad"),       # el mismo día: gana el más grave
        _r("2026-03-16", "incidente"),
        _r("2026-03-05", "sin_accidentes"),       # explícito: verde
        _r("2026-03-20", "grave"),
        _r("2026-04-01", "grave"),                 # otro mes: no cuenta aquí
    ]
    d = _cruz_calcular(regs, 2026, 3, date(2026, 3, 26), {})
    tipo = {x["dia"]: x["tipo"] for x in d["dias"]}
    assert tipo[1] == "sin_accidentes" and tipo[2] == "accidente_leve" and tipo[5] == "sin_accidentes"
    assert tipo[16] == "incidente" and tipo[20] == "grave"
    assert tipo[26] == "sin_accidentes" and tipo[27] is None and tipo[30] is None  # futuro = gris
    assert d["dias_transcurridos"] == 26 and d["mes_label"] == "Marzo 2026"


def test_resumen_cuenta_registros_y_porcentaje_sobre_dias_laborables():
    # Marzo 2026 empieza en domingo (1, 8, 15, 22 y 29 son domingos).
    regs = [_r(f"2026-03-{d:02d}", "accidente_leve") for d in (2, 3, 16, 25, 28)]
    regs += [_r("2026-03-10", "dano_propiedad"), _r("2026-03-11", "incidente")]
    d = _cruz_calcular(regs, 2026, 3, date(2026, 3, 30), {})
    assert d["dias_transcurridos"] == 30 and d["dias_laborables"] == 30 - 5  # sin los 5 domingos
    res = {r["tipo"]: r for r in d["resumen"]}
    assert res["accidente_leve"]["cantidad"] == 5 and res["accidente_leve"]["porcentaje"] == 20  # 5/25
    assert res["dano_propiedad"]["cantidad"] == 1 and res["dano_propiedad"]["porcentaje"] == 4  # 1/25
    assert res["incapacitante"]["cantidad"] == 0 and res["incapacitante"]["porcentaje"] == 0
    assert res["sin_accidentes"]["cantidad"] == 25 - 7  # días laborables sin ningún evento
    assert [r["abrev"] for r in d["resumen"]] == ["SA", "DP", "I", "PA", "INCA", "GR"]


def test_el_domingo_no_cuenta_como_dia_laborable_ni_como_dia_sin_accidentes():
    d = _cruz_calcular([], 2026, 3, date(2026, 3, 8), {})  # 1 al 8 de marzo: 2 domingos (1 y 8)
    assert d["dias_transcurridos"] == 8 and d["dias_laborables"] == 6
    sa = {r["tipo"]: r for r in d["resumen"]}["sin_accidentes"]
    assert sa["cantidad"] == 6 and sa["porcentaje"] == 100
    # Pero el domingo SÍ se pinta en la cruz (verde): la cruz muestra todos los días.
    assert [x["tipo"] for x in d["dias"] if x["dia"] in (1, 8)] == ["sin_accidentes", "sin_accidentes"]


def test_meses_pasados_completos_y_futuros_todo_gris():
    pasado = _cruz_calcular([], 2026, 1, date(2026, 3, 26), {})
    assert pasado["dias_transcurridos"] == 31 and all(x["tipo"] == "sin_accidentes" for x in pasado["dias"])
    futuro = _cruz_calcular([], 2026, 8, date(2026, 3, 26), {})
    assert futuro["dias_transcurridos"] == 0 and futuro["dias_laborables"] == 0
    assert all(x["tipo"] is None for x in futuro["dias"])
    assert all(r["porcentaje"] == 0 for r in futuro["resumen"])


def test_fuerza_de_trabajo_promedio_y_reportes_anteriores_al_formato():
    d = _cruz_calcular([{"fecha": "2026-03-04", "severidad": "moderado", "descripcion": "x"}], 2026, 3,
                       date(2026, 3, 30), {1: 200, 2: 206, 3: 0})
    assert d["fuerza_trabajo"] == {"personal": 203}  # solo personal; los días sin dato no promedian
    assert [x["tipo"] for x in d["dias"] if x["dia"] == 4] == ["incapacitante"]  # moderado → incapacitante
    assert _cruz_calcular([], 2026, 3, date(2026, 3, 30), {})["fuerza_trabajo"] == {"personal": None}


def test_bitacora_es_del_anio_sin_sin_accidentes_y_numerada():
    regs = [
        _r("2026-03-12", "dano_propiedad", trabajador="Vehículo", contratista="ALFA", ubicacion="E3 PERISUR",
           descripcion="UNIDAD PESADA DAÑA GUARNICIÓN"),
        _r("2026-03-13", "sin_accidentes"),
        _r("2026-05-22", "incidente", trabajador="Víctor Ruiz", contratista="ALFA"),
        _r("2025-12-30", "grave"),                 # otro año
    ]
    d = _cruz_calcular(regs, 2026, 5, date(2026, 5, 30), {})
    b = d["registros"]
    assert [(x["n"], x["mes"], x["abrev"]) for x in b] == [(1, "Marzo", "DP"), (2, "Mayo", "I")]
    assert b[0]["trabajador"] == "Vehículo" and b[0]["fecha"] == "2026-03-12"


def test_xlsx_replica_la_hoja_de_referencia():
    from openpyxl import load_workbook
    regs = [
        _r("2026-03-12", "dano_propiedad", trabajador="Vehículo", contratista="ALFA", ubicacion="E3 PERISUR",
           descripcion="UNIDAD PESADA DAÑA GUARNICIÓN"),
        _r("2026-04-14", "dano_propiedad", trabajador="Vehículo", contratista="ALFA", ubicacion="E3 PERISUR",
           descripcion="IMPACTA TAPIAL"),
        _r("2026-04-20", "accidente_leve", trabajador="Maria", contratista="ALFA", ubicacion="POSTE C-08"),
    ]
    data = _cruz_calcular(regs, 2026, 4, date(2026, 4, 25), {1: 203})
    ws = load_workbook(io.BytesIO(_cruz_xlsx_bytes(data, "Línea 4"))).active
    assert ws.title == "CRUZ DE SEGURIDAD" and ws["A1"].value == "CRUZ DE SEGURIDAD"
    assert ws["C4"].value == "Fuerza de Trabajo PROMEDIO" and ws["D4"].value == 203 and ws["E4"].value == "personal"
    assert ws["C5"].value == "REGISTRO DE ACCIDENTES"
    assert ws["C6"].value.startswith("SIN ACCIDENTES") and ws["D6"].value == "TIPO ACCIDENTE"
    filas = {ws.cell(r, 4).value: (ws.cell(r, 3).value, ws.cell(r, 5).value, ws.cell(r, 6).value) for r in range(7, 12)}
    # Abril 2026 hasta el 25: 25 días, 3 domingos (5, 12, 19) y 1 más (26 no llega) → 22 laborables
    assert data["dias_laborables"] == 22
    assert filas["DP"] == ("DAÑO A LA PROPIEDAD", 1, "5%") and filas["PA"][1] == 1 and filas["GR"][1] == 0
    assert ws["C7"].fill.fgColor.rgb.endswith("6DC8E8")  # color de DP en la referencia
    # La cruz: el día 20 (accidente leve) amarillo, el 26 (futuro) gris, el 1 verde.
    celdas = {ws.cell(r, c).value: ws.cell(r, c).fill.fgColor.rgb[-6:] for r in range(4, 11) for c in range(10, 19)
              if ws.cell(r, c).value}
    assert celdas[20] == "FFC000" and celdas[26] == "A6A6A6" and celdas[1] == "008040" and len(celdas) == 30
    # Bitácora: encabezado, 1ª fila de marzo sin separador, luego separador de mes en gris.
    assert [ws.cell(14, c).value for c in range(1, 9)] == [
        "Nº", "Mes", "Nombre de trabajador", "Contratista", "Tipo", "Ubicación", "Fecha", "Descripción"]
    assert [ws.cell(15, c).value for c in (1, 2, 3, 5, 7)] == [1, "Marzo", "Vehículo", "DP", "12/03/2026"]
    assert ws["A16"].value == "ABRIL" and ws["A16"].fill.fgColor.rgb.endswith("BFBFBF")
    assert [ws.cell(17, c).value for c in (1, 2, 5)] == [2, "Abril", "DP"] and ws["E18"].value == "PA"


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")
USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}


def test_endpoints_de_datos_y_excel():
    from mongomock_motor import AsyncMongoMockClient
    from fastapi import HTTPException

    db = AsyncMongoMockClient()["t"]
    old, server.db = server.db, db
    try:
        async def go():
            hoy = datetime.now(server.MX_TZ).date()
            await db.projects.insert_one({"id": "p1", "name": "Línea 4", "modulos": {}})
            await db.reports.insert_one({
                "id": "r1", "project_id": "p1", "created_at": datetime.now(timezone.utc),
                "personnel": ["3 Ayudante", "2 Cabo"], "equipment": ["1 Retroexcavadora"]})
            await db.seguridad_reportes.insert_one({
                "id": "s1", "project_id": "p1", "fecha": hoy.isoformat(), "tipo_registro": "accidente_leve",
                "trabajador": "Juan", "contratista": "Alfa", "descripcion": "Caída", "ubicacion": "E3"})
            datos = await server.get_seguridad_cruz("p1", None, USER)
            assert datos["fuerza_trabajo"] == {"personal": 5}  # 3 ayudantes + 2 cabos; la maquinaria no cuenta
            assert {r["tipo"]: r["cantidad"] for r in datos["resumen"]}["accidente_leve"] == 1
            assert datos["registros"][0]["trabajador"] == "Juan"
            resp = await server.export_seguridad_cruz_xlsx("p1", None, USER)
            blob = b"".join([c async for c in resp.body_iterator])
            assert blob[:2] == b"PK" and datos["mes"] in resp.headers["Content-Disposition"]
            with pytest.raises(HTTPException) as e:
                await server.get_seguridad_cruz("p1", "2026-13", USER)
            assert e.value.status_code == 400
        asyncio.run(go())
    finally:
        server.db = old
