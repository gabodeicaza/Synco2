"""Aspectos Sociales del reporte semanal (server.py): levantamientos (dictámenes
estructurales y fe notarial), salvamento arqueológico y pendientes de ámbito
social. Tests puros y de punta a punta del export en PDF/PPTX/DOCX contra una
Mongo simulada (`mongomock-motor`, se omite si no está instalado)."""
import asyncio
import io
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_aspectos_sociales")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import (  # noqa: E402
    _grafica_donas_png, _grafica_gantt_png, _levantamientos_bloques, _levantamientos_graficas,
    _pendientes_bloques, _salvamento_dias, _salvamento_graficas,
)

HOY = datetime.now(server.MX_TZ).date()


def _png_ok(png: bytes, ancho=1800):
    from PIL import Image
    img = Image.open(io.BytesIO(png))
    assert img.format == "PNG" and img.size[0] == ancho
    return img


# ── Levantamientos ─────────────────────────────────────────────────────────
LEV = [
    {"tipo": "dictamen_estructural", "ambito": "estacion", "total_requerido": 99, "realizados": 49},
    {"tipo": "dictamen_estructural", "ambito": "poste", "total_requerido": 193, "realizados": 19},
    {"tipo": "fe_notarial", "ambito": "estacion", "total_requerido": 99, "realizados": 0},
    {"tipo": "fe_notarial", "ambito": "poste", "total_requerido": 0, "realizados": 0},  # sin requerido: no se imprime
]


def test_levantamientos_tabla_con_porcentaje_y_omite_los_sin_requerido():
    (tabla,) = _levantamientos_bloques(LEV)
    assert tabla["rows"][0] == ["Dictámenes estructurales (estaciones)", "99", "49", "49.5%"]
    assert tabla["rows"][1][3] == "9.8%" and tabla["rows"][2] == ["Fe notarial (estaciones)", "99", "0", "0.0%"]
    assert len(tabla["rows"]) == 3 and all(len(r) == 4 for r in tabla["rows"])
    assert _levantamientos_bloques([]) == []


def test_graficas_de_dona_una_por_tipo():
    graficas = _levantamientos_graficas(LEV)
    assert [g[0] for g in graficas] == ["DICTÁMENES ESTRUCTURALES", "FE NOTARIAL"]
    for _, _, png in graficas:
        _png_ok(png)
    _png_ok(_grafica_donas_png("x", [("a", 0, 0), ("b", 5, 5)]))  # total 0 no truena


# ── Salvamento arqueológico ────────────────────────────────────────────────
def _reg(ambito, etiqueta, ini, lib, orden=0):
    return {"ambito": ambito, "etiqueta": etiqueta, "fecha_inicio": ini.isoformat(),
            "fecha_liberacion": lib.isoformat() if lib else None, "orden": orden}


def test_dias_de_salvamento():
    ini = HOY - timedelta(days=10)
    assert _salvamento_dias(_reg("poste", "S2", ini, ini + timedelta(days=4)), HOY) == 5  # inclusivo
    assert _salvamento_dias(_reg("poste", "S2", ini, None), HOY) == 11  # en curso: hasta el corte
    assert _salvamento_dias(_reg("poste", "S2", HOY + timedelta(days=3), None), HOY) is None  # aún no inicia
    assert _salvamento_dias(_reg("poste", "S2", ini, HOY + timedelta(days=9)), HOY) == 11  # liberación futura: se corta
    assert _salvamento_dias({"fecha_inicio": "basura"}, HOY) is None


def test_graficas_de_salvamento_por_ambito_y_paginadas():
    regs = [_reg("estacion", f"EST PS{i}", HOY - timedelta(days=60 - i), HOY - timedelta(days=i), i) for i in range(1, 4)]
    regs += [_reg("poste", f"S{i}", HOY - timedelta(days=100 - i), None if i % 2 else HOY - timedelta(days=i), i) for i in range(1, 31)]
    regs.append(_reg("poste", "FUTURO", HOY + timedelta(days=5), None, 99))
    graficas = _salvamento_graficas(regs, HOY)
    assert [g[0] for g in graficas] == [
        "SALVAMENTO ARQUEOLÓGICO POR ESTACIONES",
        "SALVAMENTO ARQUEOLÓGICO POR POSTES (1/2)",
        "SALVAMENTO ARQUEOLÓGICO POR POSTES (2/2)",
    ]
    for _, _, png in graficas:
        _png_ok(png)
    assert _salvamento_graficas([], HOY) == []


def test_gantt_con_una_sola_fila_no_truena():
    _png_ok(_grafica_gantt_png("t", [("S2", HOY, HOY, True, 1)], HOY))


# ── Pendientes de ámbito social ────────────────────────────────────────────
def test_pendientes_sociales_tienen_su_propio_bloque():
    regs = [
        {"contrato": "sociales", "titulo": "Liberación de S29", "descripcion": "Reunión el lunes", "orden": 1},
        {"contrato": "obra", "titulo": "Control de Avance", "orden": 1},
        {"contrato": "sociales", "titulo": "Ya liberado", "resuelto": True},
    ]
    bloques = _pendientes_bloques(regs)
    assert [b["heading"] for b in bloques] == [
        "PENDIENTES O REQUERIMIENTOS GENERALES (OBRA)", "PENDIENTES GENERALES (ASPECTOS SOCIALES)"]
    assert [i["title"] for i in bloques[1]["items"]] == ["Liberación de S29"]


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
SOCIAL = {"id": "u2", "role": server.ROLE_ESPECIALISTA, "name": "Social", "email": "s@x.mx", "project_ids": ["p1"]}
OTRO = {"id": "u3", "role": server.ROLE_ESPECIALISTA, "name": "Otro", "email": "o@x.mx", "project_ids": ["p1"]}
JEFE = {"id": "u4", "role": server.ROLE_JEFE, "name": "Jefe", "email": "j@x.mx"}
MODULOS = {"aspectos_sociales": True, "pendientes_generales": True}


def _db():
    from mongomock_motor import AsyncMongoMockClient
    return AsyncMongoMockClient()["t"]


def _con_db(coro_fn):
    db = _db()
    old = server.db
    server.db = db
    try:
        return asyncio.run(coro_fn(db))
    finally:
        server.db = old


async def _sembrar(db):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": {}, "reportes_semanales": MODULOS, "reportes_mensuales": {}},
    })
    for tipo, ambito, total, real in (("dictamen_estructural", "estacion", 99, 49), ("dictamen_estructural", "poste", 193, 19),
                                      ("fe_notarial", "estacion", 99, 0)):
        await server.set_levantamiento_social("p1", server.LevantamientoIn(tipo=tipo, ambito=ambito, total_requerido=total, realizados=real), SOCIAL)
    ini = (HOY - timedelta(days=40)).isoformat()
    await server.create_salvamento_arqueologico("p1", server.SalvamentoIn(
        ambito="estacion", etiqueta="UNIVERSIDAD PS2", fecha_inicio=ini, fecha_liberacion=(HOY - timedelta(days=30)).isoformat()), SOCIAL)
    await server.create_salvamento_arqueologico("p1", server.SalvamentoIn(ambito="poste", etiqueta="S2", fecha_inicio=ini), SOCIAL)
    await server.create_pendiente_general("p1", server.PendienteIn(
        contrato="sociales", titulo="Liberación de la S29", descripcion="Sin actividad durante esta semana"), SOCIAL)


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_semanal_imprime_aspectos_sociales(formato):
    from test_capa_financiera import _texto

    async def go(db):
        await _sembrar(db)
        common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER)
        if formato == "pdf":
            resp = await server.export_reports_pdf(
                **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None, estructura_solo=False)
        elif formato == "pptx":
            resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
        else:
            resp = await server.export_reports_docx(**common)
        return b"".join([c async for c in resp.body_iterator])
    data = _con_db(go)
    texto = _texto(formato, data)
    for esperado in (
        "ASPECTOS SOCIALES", "AVANCE EN DICTÁMENES ESTRUCTURALES Y FE NOTARIAL", "Dictámenes estructurales (estaciones)",
        "49.5%", "Fe notarial (estaciones)", "DICTÁMENES ESTRUCTURALES", "FE NOTARIAL",
        "SALVAMENTO ARQUEOLÓGICO POR ESTACIONES", "SALVAMENTO ARQUEOLÓGICO POR POSTES",
        "PENDIENTES GENERALES (ASPECTOS SOCIALES)", "Liberación de la S29",
    ):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"
    if formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 4  # 2 donas + 2 gantt


def test_modulo_apagado_no_imprime_y_no_deja_capturar():
    from test_capa_financiera import _texto

    async def go(db):
        await _sembrar(db)
        await db.projects.update_one({"id": "p1"}, {"$set": {"modulos.reportes_semanales": {}}})
        with pytest.raises(HTTPException) as e:
            await server.create_salvamento_arqueologico(
                "p1", server.SalvamentoIn(ambito="poste", etiqueta="S3", fecha_inicio=HOY.isoformat()), SOCIAL)
        assert e.value.status_code == 403
        resp = await server.export_reports_pptx(
            pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER,
            recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
        return b"".join([c async for c in resp.body_iterator])
    texto = _texto("pptx", _con_db(go))
    assert "SALVAMENTO ARQUEOLÓGICO" not in texto and "DICTÁMENES" not in texto


def test_validaciones_y_permisos():
    async def go(db):
        await _sembrar(db)
        with pytest.raises(HTTPException) as e:  # realizados > requerido
            await server.set_levantamiento_social(
                "p1", server.LevantamientoIn(tipo="fe_notarial", ambito="poste", total_requerido=5, realizados=6), SOCIAL)
        assert e.value.status_code == 400
        # upsert: reemplaza el mismo (tipo, ámbito) sin duplicar
        await server.set_levantamiento_social(
            "p1", server.LevantamientoIn(tipo="dictamen_estructural", ambito="estacion", total_requerido=99, realizados=60), SOCIAL)
        lista = await server.list_levantamientos_sociales("p1", SOCIAL)
        assert len(lista) == 3 and next(r for r in lista if r["tipo"] == "dictamen_estructural" and r["ambito"] == "estacion")["realizados"] == 60
        with pytest.raises(HTTPException):  # liberación antes del inicio
            await server.create_salvamento_arqueologico("p1", server.SalvamentoIn(
                ambito="poste", etiqueta="S9", fecha_inicio="2026-05-10", fecha_liberacion="2026-05-01"), SOCIAL)
        with pytest.raises(HTTPException):  # fecha mal formada
            await server.create_salvamento_arqueologico("p1", server.SalvamentoIn(
                ambito="poste", etiqueta="S9", fecha_inicio="31/02/2026"), SOCIAL)
        with pytest.raises(HTTPException) as e2:  # Jefe solo lee
            await server.set_levantamiento_social(
                "p1", server.LevantamientoIn(tipo="fe_notarial", ambito="poste", total_requerido=1, realizados=0), JEFE)
        assert e2.value.status_code == 403
        reg = (await server.list_salvamento_arqueologico("p1", SOCIAL))[0]
        with pytest.raises(HTTPException) as e3:  # otro especialista no borra lo ajeno
            await server.delete_salvamento_arqueologico("p1", reg["id"], OTRO)
        assert e3.value.status_code == 403
        # liberar y reabrir
        lib = await server.update_salvamento_arqueologico(
            "p1", reg["id"], server.SalvamentoPatch(fecha_liberacion=HOY.isoformat()), SOCIAL)
        assert lib["fecha_liberacion"] == HOY.isoformat()
        reab = await server.update_salvamento_arqueologico("p1", reg["id"], server.SalvamentoPatch(quitar_liberacion=True), SOCIAL)
        assert reab["fecha_liberacion"] is None
        assert (await server.delete_salvamento_arqueologico("p1", reg["id"], SOCIAL))["ok"]
    _con_db(go)
