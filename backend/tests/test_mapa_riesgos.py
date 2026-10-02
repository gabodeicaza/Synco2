"""Mapa de riesgos (server.py) y tablas largas de los reportes: numeración por
estación, qué riesgos entran en el semanal, y que las tablas del PPTX/PDF
paginen sin recortar texto ni salirse de la diapositiva (antes el PPTX usaba
filas de alto fijo, truncaba a 90 caracteres y no pasaba de diapositiva)."""
import asyncio
import io
import os
import sys
from datetime import datetime, timedelta
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_mapa_riesgos")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest  # noqa: E402

import server  # noqa: E402
from server import _pptx_fila_alto, _pptx_texto_lineas, _riesgos_grupos  # noqa: E402

HOY = datetime.now(server.MX_TZ).date()


# ── Puros ──────────────────────────────────────────────────────────────────
def _r(est, desc, ctrl="control", av=0, **extra):
    return {"estacion_nombre": est, "descripcion_riesgo": desc, "propuesta_control": ctrl, "porcentaje_avance": av, **extra}


def test_numeracion_por_estacion_quita_el_numero_tecleado_y_respeta_numerar_false():
    regs = [
        _r("E1", "1)Atropellamientos"), _r("E1", "2) Caídas"), _r("E1", "Sin número"),
        _r("Postes S-02 a S-06", "Poste 2. Riesgos 1, 2 y 6.", numerar=False),
        _r("E3", "Golpes por perforadora"), _r("E3", "3. Choque de vehículos"),
    ]
    g = {x["estacion"]: [f["descripcion"] for f in x["filas"]] for x in _riesgos_grupos(regs)}
    assert g["E1"] == ["1) Atropellamientos", "2) Caídas", "3) Sin número"]
    assert g["Postes S-02 a S-06"] == ["Poste 2. Riesgos 1, 2 y 6."]  # sin número
    assert g["E3"] == ["1) Golpes por perforadora", "2) Choque de vehículos"]  # reinicia por estación


def test_grupos_conservan_orden_avance_y_estacion_vacia():
    (g,) = _riesgos_grupos([{"descripcion_riesgo": "x", "propuesta_control": "y", "porcentaje_avance": 30}])
    assert g["estacion"] == "—" and g["filas"] == [{"descripcion": "1) x", "control": "y", "avance": 30}]
    assert _riesgos_grupos([]) == []


def test_estimacion_de_alto_de_fila_crece_con_el_texto():
    ancho = 4_000_000  # ~11 cm
    corto = _pptx_fila_alto(["30%"], [ancho], 9)
    largo = _pptx_fila_alto(["x" * 600], [ancho], 9)
    assert largo > 4 * corto
    assert _pptx_texto_lineas("a\nb\nc", ancho, 9) == 3  # los saltos de línea cuentan
    assert _pptx_texto_lineas("", ancho, 9) == 1


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
LARGO = ("Señalamientos preventivos e informativos durante los trabajos para evitar el ingreso de personas ajenas "
         "al área intervenida y abanderamiento permanente con personal capacitado. ") * 3


def _db():
    from mongomock_motor import AsyncMongoMockClient
    return AsyncMongoMockClient()["t"]


def _con_db(coro_fn):
    db = _db()
    old = server.db
    server.db = db
    try:
        return asyncio.run(coro_fn(db))
    finally:
        server.db = old


async def _proyecto(db, modulos_sem, modulos_dia=None):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": modulos_dia or {}, "reportes_semanales": modulos_sem, "reportes_mensuales": {}},
    })


async def _riesgo(db, est, desc, av, hace_dias, **extra):
    ts = datetime.now(server.timezone.utc) - timedelta(days=hace_dias)
    await db.riesgos_registros.insert_one({
        "id": f"{est}-{desc}"[:40], "project_id": "p1", "estacion_nombre": est, "descripcion_riesgo": desc,
        "propuesta_control": "control", "porcentaje_avance": av, "created_at": ts, "updated_at": ts, **extra})


async def _exportar(formato, period="week", period_type="semanal"):
    common = dict(pid="p1", period=period, period_type=period_type, area_id=None, scope=None, user=USER)
    if formato == "pdf":
        resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                               recuento_sin_datos=None, estructura_solo=False)
    elif formato == "pptx":
        resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
    else:
        resp = await server.export_reports_docx(**common)
    return b"".join([c async for c in resp.body_iterator])


def test_el_semanal_imprime_el_mapa_vigente_y_el_diario_solo_lo_del_dia():
    from test_capa_financiera import _texto

    async def go(db):
        await _proyecto(db, {"gestion_riesgos": True}, {"gestion_riesgos": True})
        await _riesgo(db, "E1", "VIGENTE antiguo al 50", 50, hace_dias=40)
        await _riesgo(db, "E1", "CERRADO antiguo al 100", 100, hace_dias=40)
        await _riesgo(db, "E3", "CERRADO esta semana", 100, hace_dias=2)
        await _riesgo(db, "E3", "NUEVO de hoy", 10, hace_dias=0)
        return (await _exportar("pptx")), (await _exportar("pptx", period="today", period_type="diario"))
    semanal, diario = _con_db(go)
    ts, td = _texto("pptx", semanal), _texto("pptx", diario)
    assert "VIGENTE antiguo al 50" in ts and "CERRADO esta semana" in ts and "NUEVO de hoy" in ts
    assert "CERRADO antiguo al 100" not in ts  # ya cerrado y fuera del período
    assert "NUEVO de hoy" in td and "VIGENTE antiguo al 50" not in td  # el diario mantiene su criterio
    # numera dentro de cada estación: E1 → 1) vigente; E3 → 1) cerrado esta semana, 2) nuevo
    assert "1) VIGENTE antiguo al 50" in ts and "1) CERRADO esta semana" in ts and "2) NUEVO de hoy" in ts


def _pptx_shapes_dentro(data: bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    fuera = [(i + 1, sh.name, sh.top + sh.height) for i, sl in enumerate(prs.slides) for sh in sl.shapes
             if sh.top is not None and sh.top + sh.height > prs.slide_height + 5_000]
    return prs, fuera


def test_mapa_largo_en_pptx_pagina_no_recorta_y_no_se_sale_de_la_diapositiva():
    from test_capa_financiera import _texto

    async def go(db):
        await _proyecto(db, {"gestion_riesgos": True})
        for e in ("E1", "E3", "E5"):
            for i in range(1, 9):
                await _riesgo(db, e, f"Riesgo {i} de {e}: " + LARGO, 10 * i, hace_dias=0)
        await _riesgo(db, "Postes S-02 a S-06", "Poste 2. Riesgos 1, 2 y 6.", 50, hace_dias=0, numerar=False)
        return await _exportar("pptx"), await _exportar("pdf")
    pptx, pdf = _con_db(go)
    prs, fuera = _pptx_shapes_dentro(pptx)
    assert not fuera, f"formas fuera de la diapositiva: {fuera[:3]}"
    texto = _texto("pptx", pptx)
    assert texto.count(LARGO.strip()[:60]) >= 24  # los 24 riesgos con su texto, sin recortar
    assert "(cont.)" in texto  # el grupo continúa en la siguiente diapositiva
    assert "Poste 2. Riesgos 1, 2 y 6." in texto and "1) Poste 2." not in texto
    assert len(prs.slides) >= 4
    from pypdf import PdfReader
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 3
    tpdf = _texto("pdf", pdf)
    assert tpdf.count(LARGO.strip()[:40]) >= 24


def test_tabla_generica_del_pptx_respeta_anchos_no_trunca_y_pagina():
    """Pruebas de Calidad (7 columnas, observaciones largas, muchas filas)."""
    from test_capa_financiera import _texto
    obs = "Se verifica concreto suministrado y se coteja la remisión contra el proyecto " * 3  # > 90 caracteres

    async def go(db):
        await _proyecto(db, {"pruebas_calidad": True})
        for i in range(1, 41):
            await server.create_prueba_calidad("p1", server.PruebaCalidadIn(
                frente=f"Frente {i}", localizacion="E1- E2", elemento="banqueta", tipo="Concreto normal",
                observaciones=obs, fecha=HOY.isoformat(), jerarquia="semanal"), {**USER, "project_ids": ["p1"]})
        return await _exportar("pptx"), await _exportar("pdf")
    pptx, pdf = _con_db(go)
    prs, fuera = _pptx_shapes_dentro(pptx)
    assert not fuera, f"formas fuera de la diapositiva: {fuera[:3]}"
    texto = _texto("pptx", pptx)
    assert texto.count(obs.strip()) == 40  # ninguna observación recortada a 90 caracteres
    assert "Frente 1" in texto and "Frente 40" in texto
    assert len(prs.slides) >= 3
    assert _texto("pdf", pdf).count("Se verifica concreto suministrado") >= 40
