"""Programa de obra en barras del reporte semanal (server.py): filas por partida
y subpartida con plazo, avance real y programado al corte, gráfica con Pillow y
export en PDF/PPTX/DOCX para obra y supervisión."""
import io
import os
import sys
from datetime import timedelta
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_programa_barras")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest  # noqa: E402

import server  # noqa: E402
from server import _grafica_programa_png, _programa_barras, _programa_barras_graficas  # noqa: E402
from test_capa_financiera import FIN, HOY, INI, USER, _concepto, _datos, _texto  # noqa: E402


def _png_ok(png: bytes):
    from PIL import Image
    img = Image.open(io.BytesIO(png))
    assert img.format == "PNG" and img.size[0] == 1800
    return img


def test_filas_por_partida_y_subpartida_con_avance_real_y_programado():
    conceptos, partidas, registros, _ = _datos()
    filas = _programa_barras(conceptos, partidas, registros, FIN)
    assert [(f["nombre"], f["nivel"]) for f in filas] == [("ESTUDIOS", 0), ("OBRA CIVIL", 0), ("Cimentación", 1)]
    est, civil, cim = filas
    assert est["real"] == pytest.approx(20.0)  # 200 / 1000
    assert civil["real"] == pytest.approx(cim["real"]) and civil["real"] == pytest.approx(400 / 3000 * 100)  # el avance futuro (999) no cuenta
    assert 40 < est["programado"] < 60  # ~mitad del plazo transcurrida al corte
    assert est["inicio"] == HOY - timedelta(days=99) and est["fin"] == HOY + timedelta(days=100)


def test_max_nivel_cero_solo_raices_y_partidas_sin_fechas_se_omiten():
    conceptos, partidas, registros, _ = _datos()
    assert [f["nombre"] for f in _programa_barras(conceptos, partidas, registros, FIN, max_nivel=0)] == ["ESTUDIOS", "OBRA CIVIL"]
    sin_fechas = [{**c, "fecha_inicio": None, "fecha_fin": None} for c in conceptos]
    assert _programa_barras(sin_fechas, partidas, registros, FIN) == []
    inactiva = [{**pa, "activo": False} if pa["id"] == "pa" else pa for pa in partidas]
    assert "ESTUDIOS" not in [f["nombre"] for f in _programa_barras(conceptos, inactiva, registros, FIN)]
    assert _programa_barras([], [], [], FIN) == []


def test_avance_se_limita_al_100_y_no_truena_con_importe_cero():
    conceptos, partidas, registros, _ = _datos()
    registros = registros + [{"concepto_id": "c1", "fecha": INI.isoformat(), "importe_ganado_congelado": "5000"}]
    assert _programa_barras(conceptos, partidas, registros, FIN)[0]["real"] == 100.0


def test_graficas_paginadas_y_png_valido():
    base = HOY - timedelta(days=50)
    partidas = [{"id": f"p{i}", "clave": str(i), "nombre": f"PARTIDA {i}", "orden": i, "activo": True} for i in range(30)]
    conceptos = [_concepto(f"c{i}", f"p{i}", 100, base, base + timedelta(days=20 + i)) for i in range(30)]
    graficas = _programa_barras_graficas(conceptos, partidas, [], INI, FIN)
    assert [g[0] for g in graficas] == ["PROGRAMA DE OBRA (1/2)", "PROGRAMA DE OBRA (2/2)"]
    assert graficas[0][1].startswith("Programa de obra de la semana del ")
    for _, _, png in graficas:
        _png_ok(png)
    sup = _programa_barras_graficas(conceptos[:2], partidas[:2], [], INI, FIN, "supervision")
    assert sup[0][0] == "PROGRAMA DE SUPERVISIÓN" and "servicios de supervisión" in sup[0][1]
    assert _programa_barras_graficas([], [], [], INI, FIN) == []


def test_una_sola_fila_y_un_dia_no_truenan():
    fila = {"nombre": "X", "nivel": 0, "inicio": HOY, "fin": HOY, "real": 0.0, "programado": 0.0}
    _png_ok(_grafica_programa_png("t", [fila], HOY))


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_imprime_el_programa_de_obra_y_de_supervision(formato):
    pytest.importorskip("mongomock_motor")
    from test_contrato_financiero import _con_db, _sembrar_con_supervision

    async def go(db):
        await _sembrar_con_supervision(db)
        common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER)
        if formato == "pdf":
            resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                   recuento_sin_datos=None, estructura_solo=False)
        elif formato == "pptx":
            resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
        else:
            resp = await server.export_reports_docx(**common)
        return b"".join([c async for c in resp.body_iterator])
    data = _con_db(go)
    texto = _texto(formato, data)
    for esperado in ("PROGRAMA DE OBRA", "PROGRAMA DE SUPERVISIÓN", "Programa de obra de la semana del",
                     "Programa de servicios de supervisión de la semana del", "TABLA DE AVANCE POR PARTIDAS (OBRA)"):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"
    if formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 6  # 2 programas + 2×2 gráficas
        assert not [1 for sl in prs.slides for sh in sl.shapes if sh.top + sh.height > prs.slide_height + 5_000]
