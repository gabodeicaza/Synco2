"""Control de Calidad con fotos y Concentrado de Pruebas (server.py): tests puros
y de punta a punta del export semanal en PDF, PPTX y DOCX contra una Mongo
simulada (`mongomock-motor`, se omite si no está instalado)."""
import asyncio
import io
import os
import sys
from datetime import datetime
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_control_calidad")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import _control_calidad_bloques, _foto_comp_desc, _foto_comp_etiqueta  # noqa: E402

_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


# ── Puros ──────────────────────────────────────────────────────────────────
def test_etiqueta_vacia_explicita_no_cae_a_original_actual():
    assert _foto_comp_etiqueta({"etiqueta_actual": ""}, "actual", "ACTUAL") == ""
    assert _foto_comp_etiqueta({}, "actual", "ACTUAL") == "ACTUAL"
    assert _foto_comp_etiqueta({"fecha_original": "2026-06-26"}, "original", "ORIGINAL") == "ORIGINAL · 26/06/2026"
    assert _foto_comp_etiqueta({"etiqueta_original": "TAPIAL"}, "original", "ORIGINAL") == "TAPIAL"
    assert _foto_comp_desc({"desc_original": "  texto  "}, "original") == "texto"
    assert _foto_comp_desc({}, "actual") == ""


def test_bloques_emparejan_de_dos_en_dos_y_omiten_registros_sin_foto():
    regs = [
        {"node_id": "n1", "foto": "A", "descripcion": "da", "pie": "pa", "fecha": "2026-06-24"},
        {"node_id": "n1", "foto": "B", "descripcion": "db", "pie": "pb", "fecha": "2026-06-25"},
        {"node_id": "n1", "foto": "C", "descripcion": "dc", "fecha": "2026-06-26"},
        {"node_id": "n1", "foto": None, "descripcion": "sin foto", "fecha": "2026-06-27"},
        {"node_id": None, "grupo": "Suministro Gefsa-Puebla", "foto": "D", "fecha": "2026-06-27"},
    ]
    bloques = _control_calidad_bloques(regs, {"n1": "Estación 1 Universidad"})
    assert len(bloques) == 2 and all(b["type"] == "fotos_comparativa" for b in bloques)
    it1, it2 = bloques[0]["items"]
    assert it1["titulo"] == "CONTROL DE CALIDAD · ESTACIÓN 1 UNIVERSIDAD"
    assert (it1["original"], it1["actual"]) == ("A", "B")
    assert (it1["desc_original"], it1["etiqueta_actual"]) == ("da", "pb")
    assert it1["desc_encabezado"] == "ACTIVIDADES DE CONTROL DE CALIDAD"
    assert (it2["original"], it2["actual"], it2["etiqueta_actual"]) == ("C", None, "")
    assert it2["sin_placeholder"] is True
    assert bloques[1]["items"][0]["titulo"].endswith("SUMINISTRO GEFSA-PUEBLA")


def test_sin_fotos_no_hay_bloques():
    assert _control_calidad_bloques([], {}) == []
    assert _control_calidad_bloques([{"foto": None}], {}) == []


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
ESPEC = {"id": "u2", "role": server.ROLE_ESPECIALISTA, "name": "Calidad", "email": "q@x.mx", "project_ids": ["p1"]}


def _db():
    from mongomock_motor import AsyncMongoMockClient
    return AsyncMongoMockClient()["t"]


async def _proyecto(db, modulos):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": {}, "reportes_semanales": modulos, "reportes_mensuales": {}},
    })
    await db.location_nodes.insert_one({"id": "n1", "project_id": "p1", "name": "ESTACIÓN 1 UNIVERSIDAD",
                                        "parent_id": None, "depth": 1, "order": 0, "is_leaf": True})


async def _capturar():
    hoy = datetime.now(server.MX_TZ).date().isoformat()
    for desc, pie in (("Verificación de hincado de pila", "Se verifica proceso de habilitado"),
                      ("Verificación de suministro Poste S-2", "Vista de suministro")):
        await server.create_control_calidad("p1", server.ControlCalidadIn(
            image=_PNG, node_id="n1", descripcion=desc, pie=pie, fecha=hoy, jerarquia="semanal"), ESPEC)
    await server.create_control_calidad("p1", server.ControlCalidadIn(
        image=_PNG, grupo="Suministro Gefsa-Puebla", descripcion="Cotejo de suministro", pie="Plataforma", fecha=hoy, jerarquia="semanal"), ESPEC)
    await server.create_prueba_calidad("p1", server.PruebaCalidadIn(
        node_id="n1", frente="Banqueta adecuación Poste S-2", localizacion="E1- E2", elemento="banqueta",
        tipo="Concreto normal", observaciones="Se verifica concreto suministrado", fecha=hoy, jerarquia="semanal"), ESPEC)
    # Fuera del periodo: no debe aparecer.
    await server.create_prueba_calidad("p1", server.PruebaCalidadIn(
        frente="PRUEBA ANTIGUA", fecha="2020-01-01", jerarquia="semanal"), ESPEC)


def _exportar(formato, modulos, apagar_antes_de_exportar=False):
    db = _db()
    old = server.db
    server.db = db
    try:
        async def go():
            await _proyecto(db, modulos)
            await _capturar()
            if apagar_antes_de_exportar:
                await db.projects.update_one({"id": "p1"}, {"$set": {"modulos.reportes_semanales": {}}})
            common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER)
            if formato == "pdf":
                resp = await server.export_reports_pdf(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None,
                    estructura_solo=False)
            elif formato == "pptx":
                resp = await server.export_reports_pptx(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
            else:
                resp = await server.export_reports_docx(**common)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def _texto(formato, data):
    if formato == "pdf":
        from pypdf import PdfReader
        t = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    elif formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        partes = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    partes.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    partes += [c.text for r in sh.table.rows for c in r.cells]
        t = "\n".join(partes)
    else:
        from docx import Document
        d = Document(io.BytesIO(data))
        t = "\n".join([p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells])
    return " ".join(t.split())


TODOS = {"control_calidad": True, "pruebas_calidad": True}


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_semanal_imprime_control_y_pruebas(formato):
    texto = _texto(formato, _exportar(formato, TODOS))
    for esperado in (
        "CONTROL DE CALIDAD", "ESTACIÓN 1 UNIVERSIDAD", "ACTIVIDADES DE CONTROL DE CALIDAD",
        "Verificación de hincado de pila", "Se verifica proceso de habilitado", "Vista de suministro",
        "SUMINISTRO GEFSA-PUEBLA", "Cotejo de suministro",
        "CONCENTRADO DE PRUEBAS DE CALIDAD", "Banqueta adecuación Poste S-2", "E1- E2", "No Aplica",
        "Se verifica concreto suministrado",
    ):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"
    assert "PRUEBA ANTIGUA" not in texto
    assert "ORIGINAL" not in texto and "ACTUAL" not in texto.replace("ACTUALIZ", "")


def test_fotos_en_pptx_y_modulo_apagado_no_imprime():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx", TODOS)))
    assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 3
    texto = _texto("pptx", _exportar("pptx", TODOS, apagar_antes_de_exportar=True))
    assert "ACTIVIDADES DE CONTROL DE CALIDAD" not in texto and "Banqueta adecuación" not in texto


def test_permisos_y_validaciones():
    db = _db()
    old = server.db
    server.db = db
    try:
        async def go():
            await _proyecto(db, TODOS)
            with pytest.raises(HTTPException) as e:
                await server.create_control_calidad("p1", server.ControlCalidadIn(image="no-es-imagen"), ESPEC)
            assert e.value.status_code == 400
            with pytest.raises(HTTPException):
                await server.create_prueba_calidad("p1", server.PruebaCalidadIn(frente="  "), ESPEC)
            with pytest.raises(HTTPException):
                await server.create_prueba_calidad("p1", server.PruebaCalidadIn(frente="x", node_id="otro"), ESPEC)
            propio = await server.create_prueba_calidad("p1", server.PruebaCalidadIn(frente="mio"), ESPEC)
            ajeno = {"id": "u3", "role": server.ROLE_ESPECIALISTA, "name": "Otro", "email": "o@x.mx", "project_ids": ["p1"]}
            with pytest.raises(HTTPException) as e2:
                await server.delete_prueba_calidad("p1", propio["id"], ajeno)
            assert e2.value.status_code == 403
            assert (await server.delete_prueba_calidad("p1", propio["id"], ESPEC))["ok"]
            jefe = {"id": "u4", "role": server.ROLE_JEFE, "name": "Jefe", "email": "j@x.mx"}
            with pytest.raises(HTTPException) as e3:
                await server.create_prueba_calidad("p1", server.PruebaCalidadIn(frente="x"), jefe)
            assert e3.value.status_code == 403
        asyncio.run(go())
    finally:
        server.db = old
