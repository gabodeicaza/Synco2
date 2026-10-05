"""Creador de módulos (server.py, `_cm_*`): fórmulas, saneamiento de la
definición, captura de registros, bloques de informe por modo (texto / tabla /
resumen / fotos / oculto) y exportación PDF/PPTX/DOCX contra una Mongo
simulada (`mongomock-motor`, se omite si no está instalado)."""
import asyncio
import io
import os
import sys
from datetime import datetime
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_modulos_personalizados")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402

_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


# ── Fórmulas ───────────────────────────────────────────────────────────────
def test_formula_precedencia_parentesis_y_vacios():
    ev = server._cm_formula_eval
    assert ev("{a} + {b} * 2", {"a": 1, "b": 3}) == 7
    assert ev("({a} + {b}) * 2", {"a": 1, "b": 3}) == 8
    assert ev("-{a} + 10", {"a": 4}) == 6
    assert ev("{a} / {b}", {"a": 1, "b": 0}) is None  # división entre cero
    assert ev("{a} + {falta}", {"a": 5}) == 5  # referencia vacía = 0


@pytest.mark.parametrize("mala", ["", "{a} +", "(1 + 2", "1 2", "__import__('os')", "{a} ** 2 ; x"])
def test_formula_invalida_no_ejecuta_codigo(mala):
    with pytest.raises(ValueError):
        server._cm_formula_eval(mala, {"a": 1})


def test_numeros_tecleados_a_mano():
    n = server._cm_num
    assert n("$1,234.50") == 1234.5 and n("12%") == 12 and n("3,5") == 3.5 and n("  ") is None and n(None) is None
    for malo in ("abc", True, float("nan")):
        with pytest.raises((ValueError, TypeError)):
            n(malo)


# ── Saneamiento de la definición ───────────────────────────────────────────
def _campos():
    return [
        {"id": "desc", "etiqueta": "Descripción", "tipo": "texto", "requerido": True},
        {"id": "cant", "etiqueta": "Cantidad", "tipo": "numero", "unidad": "m3", "agregacion": "suma"},
        {"id": "pu", "etiqueta": "P.U.", "tipo": "moneda"},
        {"id": "imp", "etiqueta": "Importe", "tipo": "calculado", "formato": "moneda", "formula": "{cant} * {pu}", "agregacion": "suma"},
        {"id": "ok", "etiqueta": "Conforme", "tipo": "interruptor", "agregacion": "conteo"},
        {"id": "motivo", "etiqueta": "Motivo", "tipo": "texto", "visible_si": {"campo": "ok", "igual": False}},
        {"id": "tipo", "etiqueta": "Tipo", "tipo": "seleccion", "opciones": ["A", "B", "A", " "]},
        {"id": "foto", "etiqueta": "Evidencia", "tipo": "foto"},
    ]


def test_campos_validos_se_normalizan():
    out = server._cm_sanitize_campos(_campos())
    assert [c["id"] for c in out] == ["desc", "cant", "pu", "imp", "ok", "motivo", "tipo", "foto"]
    assert out[6]["opciones"] == ["A", "B"]
    assert out[3]["agregacion"] == "suma" and out[0]["requerido"] is True
    assert out[5]["visible_si"] == {"campo": "ok", "igual": False}


@pytest.mark.parametrize("cambio,fragmento", [
    ({"etiqueta": "  "}, "etiqueta"),
    ({"tipo": "inventado"}, "Tipo"),
])
def test_campos_invalidos(cambio, fragmento):
    with pytest.raises(HTTPException) as e:
        server._cm_sanitize_campos([{"etiqueta": "x", "tipo": "texto", **cambio}])
    assert e.value.status_code == 400 and fragmento in e.value.detail


def test_seleccion_sin_opciones_y_formula_mala_se_rechazan():
    with pytest.raises(HTTPException):
        server._cm_sanitize_campos([{"etiqueta": "S", "tipo": "seleccion", "opciones": []}])
    with pytest.raises(HTTPException):
        server._cm_sanitize_campos([{"etiqueta": "N", "tipo": "numero", "id": "n"},
                                    {"etiqueta": "C", "tipo": "calculado", "formula": "{n} +"}])
    with pytest.raises(HTTPException):  # una fórmula no puede apuntar a un campo de texto
        server._cm_sanitize_campos([{"etiqueta": "T", "tipo": "texto", "id": "t"},
                                    {"etiqueta": "C", "tipo": "calculado", "formula": "{t} + 1"}])


def test_ids_repetidos_o_raros_se_regeneran():
    out = server._cm_sanitize_campos([
        {"id": "x", "etiqueta": "A", "tipo": "texto"}, {"id": "x", "etiqueta": "B", "tipo": "texto"},
        {"id": "Mal Id!", "etiqueta": "C", "tipo": "texto"},
    ])
    assert len({c["id"] for c in out}) == 3


def test_salida_y_opciones_por_defecto():
    s = server._cm_sanitize_salida({"reportes_semanales": {"modo": "tabla", "periodo": "acumulado"}, "basura": 1})
    assert s["reportes_semanales"] == {"modo": "tabla", "periodo": "acumulado"}
    assert s["reportes_diarios"] == {"modo": "oculto", "periodo": "periodo"}
    o = server._cm_sanitize_opciones({"limite": 9999, "agrupar_por": "no_existe"}, [])
    assert o["limite"] == 0 and o["agrupar_por"] is None and o["mostrar_fecha"] is True and o["imprimir_vacio"] is False


# ── Valores y registros ────────────────────────────────────────────────────
def _modulo(**extra):
    campos = server._cm_sanitize_campos(_campos())
    return {"id": "m1", "nombre": "Control de materiales", "campos": campos, "activo": True,
            "salida": server._cm_sanitize_salida({"reportes_semanales": {"modo": "tabla"}}),
            "opciones": server._cm_sanitize_opciones(extra.pop("opciones", {}), campos), **extra}


def test_registro_calcula_oculta_y_valida():
    async def go():
        m = _modulo()
        valores, fotos, _ = await server._cm_construye_registro(
            m, {"valores": {"desc": "Acero", "cant": "10", "pu": "$25.50", "ok": True, "motivo": "no debería guardarse"}}, None)
        assert valores["imp"] == 255 and "motivo" not in valores  # calculado + visible_si
        valores2, _, _ = await server._cm_construye_registro(
            m, {"valores": {"desc": "Concreto", "ok": False, "motivo": "falló prueba"}}, None)
        assert valores2["motivo"] == "falló prueba" and valores2["imp"] == 0
        with pytest.raises(HTTPException) as e:
            await server._cm_construye_registro(m, {"valores": {"cant": 1}}, None)
        assert "Descripción" in e.value.detail  # requerido
        with pytest.raises(HTTPException):
            await server._cm_construye_registro(m, {"valores": {"desc": "x", "cant": "abc"}}, None)
        with pytest.raises(HTTPException):
            await server._cm_construye_registro(m, {"valores": {"desc": "x", "tipo": "Z"}}, None)
        _, fotos, thumbs = await server._cm_construye_registro(m, {"valores": {"desc": "x", "foto": _PNG}}, None)
        assert fotos["foto"].startswith("data:image/jpeg") and "foto" in thumbs
        with pytest.raises(HTTPException):
            await server._cm_construye_registro(m, {"valores": {"desc": "x", "foto": "no-es-imagen"}}, None)
    asyncio.run(go())


def test_edicion_conserva_lo_que_no_viene():
    async def go():
        m = _modulo()
        previo = {"valores": {"desc": "Acero", "cant": 2.0, "pu": 3.0}, "fotos": {}, "thumbs": {}}
        valores, _, _ = await server._cm_construye_registro(m, {"valores": {"cant": 5}}, previo)
        assert valores["desc"] == "Acero" and valores["imp"] == 15
    asyncio.run(go())


# ── Bloques de informe ─────────────────────────────────────────────────────
def _regs():
    base = [("Acero", 10, 25.5, True, "2026-06-24"), ("Concreto", 4, 100, False, "2026-06-25"), ("Grava", 6, 10, True, "2026-06-26")]
    out = []
    for i, (d, c, p, ok, f) in enumerate(base):
        out.append({"id": f"r{i}", "fecha": f, "created_by_name": "Ing. Pérez", "created_at": datetime(2026, 6, 24 + i),
                    "valores": {"desc": d, "cant": c, "pu": p, "imp": c * p, "ok": ok, "tipo": "A"}})
    return out


def test_modo_tabla_con_totales_y_formato():
    m = _modulo()
    sec = server._cm_construye_seccion(m, _regs(), "reportes_semanales")
    assert sec["titulo"] == "CONTROL DE MATERIALES"
    t = sec["blocks"][0]
    assert t["type"] == "table"
    assert t["headers"] == ["FECHA", "DESCRIPCIÓN", "CANTIDAD", "P.U.", "IMPORTE", "CONFORME", "MOTIVO", "TIPO"]
    assert t["rows"][0][:6] == ["24/06/2026", "Acero", "10 m3", "$25.50", "$255.00", "Sí"]
    assert t["rows"][1][5] == "No"
    total = t["rows"][-1]
    assert total[2] == "Total: 20 m3" and total[4] == "Total: $715.00" and total[5] == "Cuenta: 2"
    assert abs(sum(t["col_widths"]) - 1) < 0.01


def test_apagar_totales_y_numerar():
    m = _modulo(opciones={"mostrar_totales": False, "numerar": True, "mostrar_fecha": False})
    t = server._cm_construye_seccion(m, _regs(), "reportes_semanales")["blocks"][0]
    assert t["headers"][0] == "No." and len(t["rows"]) == 3 and t["rows"][2][0] == "3"


def test_oculto_y_vacio_no_imprimen():
    m = _modulo()
    assert server._cm_construye_seccion(m, _regs(), "reportes_diarios") is None  # oculto
    assert server._cm_construye_seccion(m, [], "reportes_semanales") is None  # sin datos: nada
    m2 = _modulo(opciones={"imprimir_vacio": True, "texto_vacio": "Nada que reportar"})
    sec = server._cm_construye_seccion(m2, [], "reportes_semanales")
    assert sec and sec["blocks"][-1]["text"] == "Nada que reportar" or sec["blocks"][0]["text"] == "Nada que reportar"


def test_modo_texto_con_y_sin_plantilla():
    m = _modulo(salida=server._cm_sanitize_salida({"reportes_semanales": {"modo": "texto"}}))
    items = server._cm_construye_seccion(m, _regs(), "reportes_semanales")["blocks"][0]["items"]
    assert items[0]["title"] == "Acero" and "Cantidad: 10 m3" in items[0]["body"] and items[0]["meta"] == "24/06/2026"
    mp = _modulo(salida=m["salida"], opciones={"plantilla_texto": "Se recibió {cantidad} de {descripcion} ({CONFORME})",
                                               "mostrar_autor": True})
    it = server._cm_construye_seccion(mp, _regs(), "reportes_semanales")["blocks"][0]["items"][0]
    assert it["body"] == "Se recibió 10 m3 de Acero (Sí)" and "Ing. Pérez" in it["meta"]


def test_modo_resumen_orden_limite_y_ultimo():
    sal = server._cm_sanitize_salida({"reportes_semanales": {"modo": "resumen"}})
    sec = server._cm_construye_seccion(_modulo(salida=sal), _regs(), "reportes_semanales")
    assert sec["blocks"][0]["text"] == "3 registros en el período"
    filas = dict(map(tuple, sec["blocks"][1]["rows"]))
    assert filas["Total · Cantidad"] == "20 m3" and filas["Cuenta · Conforme"] == "2"
    desc = _modulo(opciones={"orden": "desc", "limite": 2})
    t = server._cm_construye_seccion(desc, _regs(), "reportes_semanales")["blocks"][0]
    assert [r[1] for r in t["rows"][:2]] == ["Grava", "Concreto"]
    ult = _modulo(salida=server._cm_sanitize_salida({"reportes_semanales": {"modo": "tabla", "periodo": "ultimo"}}))
    t = server._cm_construye_seccion(ult, _regs(), "reportes_semanales")["blocks"][0]
    assert t["rows"][0][1] == "Grava" and len([r for r in t["rows"] if r[0] != "TOTALES"]) == 1


def test_agrupar_por_genera_un_bloque_por_grupo():
    regs = _regs()
    regs[1]["valores"]["tipo"] = "B"
    m = _modulo(opciones={"agrupar_por": "tipo"})
    bloques = server._cm_construye_seccion(m, regs, "reportes_semanales")["blocks"]
    assert [b["heading"] for b in bloques] == ["TIPO: A", "TIPO: B"]


def test_modo_fotos_y_sin_fotos():
    sal = server._cm_sanitize_salida({"reportes_semanales": {"modo": "fotos"}})
    m = _modulo(salida=sal)
    regs = _regs()
    assert server._cm_construye_seccion(m, regs, "reportes_semanales") is None  # nada con foto
    for r in regs:
        r["fotos"] = {"foto": _PNG}
    sec = server._cm_construye_seccion(m, regs, "reportes_semanales")
    items = sec["blocks"][0]["items"]
    assert sec["blocks"][0]["type"] == "fotos_comparativa" and len(items) == 2
    assert items[0]["desc_original"] == "Acero" and items[1]["actual"] is None


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

COORD = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
ESPEC = {"id": "u2", "role": server.ROLE_ESPECIALISTA, "name": "Calidad", "email": "q@x.mx", "project_ids": ["p1"]}
JEFE = {"id": "u4", "role": server.ROLE_JEFE, "name": "Jefe", "email": "j@x.mx"}


def _db():
    from mongomock_motor import AsyncMongoMockClient
    return AsyncMongoMockClient()["t"]


async def _proyecto(db):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": {}, "reportes_semanales": {}, "reportes_mensuales": {}},
    })


async def _crear_modulo(modo_semanal="tabla", **extra):
    body = {
        "nombre": "Control de materiales", "icono": "cube-outline", "campos": _campos(),
        "salida": {"reportes_semanales": {"modo": modo_semanal}}, **extra,
    }
    return await server.create_custom_module("p1", body, COORD)


def _con_db(fn):
    def wrapper():
        db = _db()
        old = server.db
        server.db = db
        try:
            async def go():
                await _proyecto(db)
                await fn(db)
            asyncio.run(go())
        finally:
            server.db = old
    return wrapper


@_con_db
async def test_crud_permisos_y_borrado_confirmado(db):
    m = await _crear_modulo()
    assert m["key"] == f"custom_{m['id']}" and m["orden"] == 0
    # El Coordinador ve todo; el especialista solo lo activo
    assert len(await server.list_custom_modules("p1", todos=True, user=COORD)) == 1
    hoy = datetime.now(server.MX_TZ).date().isoformat()
    r = await server.create_custom_module_record(
        "p1", m["id"], {"fecha": hoy, "valores": {"desc": "Acero", "cant": 3, "pu": 2, "foto": _PNG}}, ESPEC)
    assert r["valores"]["imp"] == 6 and r["tiene_fotos"] == ["foto"] and "fotos" not in r
    lista = await server.list_custom_module_records("p1", m["id"], user=ESPEC)
    assert len(lista) == 1 and "fotos" not in lista[0]
    completo = await server.get_custom_module_record("p1", m["id"], r["id"], ESPEC)
    assert completo["fotos"]["foto"].startswith("data:image/jpeg")
    # Jefe (solo lectura) no captura; un tercero no edita lo ajeno
    with pytest.raises(HTTPException) as e:
        await server.create_custom_module_record("p1", m["id"], {"valores": {"desc": "x"}}, JEFE)
    assert e.value.status_code == 403
    ajeno = {"id": "u3", "role": server.ROLE_ESPECIALISTA, "name": "Otro", "email": "o@x.mx", "project_ids": ["p1"]}
    with pytest.raises(HTTPException) as e2:
        await server.update_custom_module_record("p1", m["id"], r["id"], {"valores": {"cant": 9}}, ajeno)
    assert e2.value.status_code == 403
    upd = await server.update_custom_module_record("p1", m["id"], r["id"], {"valores": {"cant": 5}}, ESPEC)
    assert upd["valores"]["imp"] == 10
    # Módulo solo para coordinación
    await server.update_custom_module("p1", m["id"], {"solo_coordinacion": True}, COORD)
    with pytest.raises(HTTPException) as e3:
        await server.create_custom_module_record("p1", m["id"], {"valores": {"desc": "x"}}, ESPEC)
    assert e3.value.status_code == 403
    # Apagado: el especialista ya no lo ve ni captura
    await server.update_custom_module("p1", m["id"], {"activo": False}, COORD)
    assert await server.list_custom_modules("p1", todos=True, user=ESPEC) == []
    # Cambiar campos conserva los registros; borrar con datos exige confirmación
    with pytest.raises(HTTPException) as e4:
        await server.delete_custom_module("p1", m["id"], confirmar=False, user=COORD)
    assert e4.value.status_code == 409
    res = await server.delete_custom_module("p1", m["id"], confirmar=True, user=COORD)
    assert res["registros_borrados"] == 1 and await db.custom_modules.count_documents({}) == 0


@_con_db
async def test_duplicar_y_ordenar(db):
    a = await _crear_modulo()
    b = await _crear_modulo()
    assert (a["orden"], b["orden"]) == (0, 1)
    copia = await server.duplicate_custom_module("p1", a["id"], {}, COORD)
    assert copia["nombre"].endswith("(copia)") and copia["activo"] is False and copia["id"] != a["id"]
    await server.reorder_custom_modules("p1", {"ids": [b["id"], a["id"]]}, COORD)
    nombres = [m["id"] for m in await server.list_custom_modules("p1", todos=True, user=COORD)]
    assert nombres[:2] == [b["id"], a["id"]]


@_con_db
async def test_vista_previa_con_muestra_y_sin_guardar(db):
    m = await _crear_modulo()
    out = await server.preview_custom_module("p1", {"module_id": m["id"], "categoria": "reportes_semanales"}, COORD)
    assert out["muestra"] is True and out["seccion"]["blocks"][0]["type"] == "table"
    # Cambiar el modo en el borrador (sin guardar) cambia la vista previa
    out2 = await server.preview_custom_module("p1", {
        "module_id": m["id"], "categoria": "reportes_semanales",
        "definicion": {"salida": {"reportes_semanales": {"modo": "texto"}}}}, COORD)
    assert out2["seccion"]["blocks"][0]["type"] == "list"
    guardado = await server.get_custom_module("p1", m["id"], COORD)
    assert guardado["salida"]["reportes_semanales"]["modo"] == "tabla"
    out3 = await server.preview_custom_module("p1", {"categoria": "reportes_diarios", "definicion": {"nombre": "Nuevo", "campos": _campos()}}, COORD)
    assert out3["seccion"] is None and out3["modo"] == "oculto"


def _exportar(formato, modo_semanal="tabla", modulo_activo=True, categoria_periodo="week"):
    db = _db()
    old = server.db
    server.db = db
    try:
        async def go():
            await _proyecto(db)
            m = await _crear_modulo(modo_semanal=modo_semanal,
                                    opciones={"incluir_fotos": modo_semanal in ("texto", "tabla")})
            hoy = datetime.now(server.MX_TZ).date().isoformat()
            for d, c, p in (("Acero corrugado", 10, 25.5), ("Concreto premezclado", 4, 100)):
                await server.create_custom_module_record(
                    "p1", m["id"], {"fecha": hoy, "valores": {"desc": d, "cant": c, "pu": p, "foto": _PNG}}, ESPEC)
            await server.create_custom_module_record(
                "p1", m["id"], {"fecha": "2020-01-01", "valores": {"desc": "REGISTRO ANTIGUO"}}, ESPEC)
            if not modulo_activo:
                await server.update_custom_module("p1", m["id"], {"activo": False}, COORD)
            common = dict(pid="p1", period=categoria_periodo, period_type="semanal", area_id=None, scope=None, user=COORD)
            if formato == "pdf":
                resp = await server.export_reports_pdf(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None, estructura_solo=False)
            elif formato == "pptx":
                resp = await server.export_reports_pptx(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
            else:
                resp = await server.export_reports_docx(**common)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def _texto(formato, data):
    if formato == "pdf":
        from pypdf import PdfReader
        t = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    elif formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        partes = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    partes.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    partes += [c.text for r in sh.table.rows for c in r.cells]
        t = "\n".join(partes)
    else:
        from docx import Document
        d = Document(io.BytesIO(data))
        t = "\n".join([p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells])
    return " ".join(t.split())


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_semanal_imprime_el_modulo_personalizado(formato):
    texto = _texto(formato, _exportar(formato))
    for esperado in ("CONTROL DE MATERIALES", "Acero corrugado", "Concreto premezclado", "$255.00", "Total: 14 m3"):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"
    assert "REGISTRO ANTIGUO" not in texto  # fuera del período


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_modo_texto_en_los_tres_formatos(formato):
    texto = _texto(formato, _exportar(formato, modo_semanal="texto"))
    assert "Acero corrugado" in texto and "Cantidad: 10 m3" in texto


def test_oculto_o_apagado_no_se_imprime():
    for kwargs in ({"modo_semanal": "oculto"}, {"modulo_activo": False}):
        texto = _texto("docx", _exportar("docx", **kwargs))
        assert "Acero corrugado" not in texto and "CONTROL DE MATERIALES" not in texto


def test_fotos_en_pptx():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx", modo_semanal="fotos")))
    assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 2
