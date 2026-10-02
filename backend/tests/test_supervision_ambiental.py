"""Módulo Supervisión Ambiental (todo editable): informes con % y fecha de cierre,
autorizaciones, arbolado autorizado vs ejecutado por autorización, resarcimiento
con fotos, tabla resumen de estudios; CRUD, validaciones y export PDF/DOCX/PPTX."""
import asyncio
import io
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_ambiental")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import ROLE_COORD, ROLE_JEFE, ROLE_SUB, _fecha_larga_es  # noqa: E402

pytest.importorskip("mongomock_motor")
_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
COORD = {"id": "u1", "role": ROLE_COORD, "name": "C", "email": "c@x.mx", "project_ids": ["p1"]}
SUB = {"id": "u2", "role": ROLE_SUB, "name": "S", "email": "s@x.mx", "project_ids": ["p1"]}
JEFE = {"id": "u3", "role": ROLE_JEFE, "name": "J", "email": "j@x.mx", "project_ids": ["p1"]}
ACTIVO = {"reportes_diarios": {}, "reportes_mensuales": {}, "reportes_semanales": {"supervision_ambiental": True}}


def _con_db(fn, modulos=ACTIVO):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await db.projects.insert_one({"id": "p1", "name": "Proyecto Demo", "color_institucional": "#7A1030", "modulos": modulos})
            return await fn(db)
        return asyncio.run(go())
    finally:
        server.db = old


def test_fecha_larga():
    assert _fecha_larga_es("2026-09-17") == "17 de septiembre de 2026"


async def _sembrar_todo():
    c = server.create_ambiental
    await server.set_ambiental_config("p1", server.AmbientalConfigIn(
        referencia="Anexo 1 del Acuerdo y Resolución Admva. 05134/2026", total_a_restituir=46), COORD)
    await c("p1", server.AmbientalIn(seccion="informe", nombre="3er Informe Trimestral", avance=98, fecha_cierre="17/09/2026"), COORD)
    await c("p1", server.AmbientalIn(seccion="informe", nombre="1er Informe Semestral", avance=5, fecha_cierre="2026-12-02"), COORD)
    await c("p1", server.AmbientalIn(seccion="autorizacion", tipo="SEDEMA", folio="SEDEMA/DGEIRA/DEIAR/05134/2026"), COORD)
    await c("p1", server.AmbientalIn(seccion="arbolado", nombre="MIA-E", poda_aut=143, trasplante_aut=231, derribo_aut=17,
                                     poda_ej=39, trasplante_ej=179, derribo_ej=17), COORD)
    await c("p1", server.AmbientalIn(seccion="arbolado", nombre="Nueva afectación", poda_aut=14, trasplante_aut=39, derribo_aut=5,
                                     poda_ej=0, trasplante_ej=6, derribo_ej=3), COORD)
    await c("p1", server.AmbientalIn(seccion="resarcimiento", lugar="Plantaciones en Parque Cantera", cantidad=30,
                                     pie="Se plantan 30 ejemplares", image=_PNG), COORD)
    await c("p1", server.AmbientalIn(seccion="resarcimiento", lugar="Plantaciones en Anillo Circunvalación", cantidad=16, image=_PNG), COORD)
    await c("p1", server.AmbientalIn(seccion="avance_estudio", actividad="Dictamen técnico forestal", avance=100), COORD)


def test_crud_normaliza_y_es_editable():
    async def go(db):
        a = await server.create_ambiental("p1", server.AmbientalIn(seccion="informe", nombre=" 3er Informe ", avance=98, fecha_cierre="17/09/2026"), COORD)
        assert a["nombre"] == "3er Informe" and a["fecha_cierre"] == "2026-09-17" and a["orden"] == 1
        b = await server.create_ambiental("p1", server.AmbientalIn(seccion="informe", nombre="Otro", avance=5), SUB)
        assert b["orden"] == 2
        # editable: cambia avance, quita la fecha y renombra
        u = await server.update_ambiental("p1", a["id"], server.AmbientalPatch(avance=100, fecha_cierre="", nombre="3er Informe Trimestral"), COORD)
        assert u["avance"] == 100 and u["fecha_cierre"] is None and u["nombre"] == "3er Informe Trimestral"
        listado = await server.list_ambiental("p1", COORD)
        assert [i["nombre"] for i in listado["items"]] == ["3er Informe Trimestral", "Otro"]
        assert all("foto" not in i for i in listado["items"])
        await server.delete_ambiental("p1", b["id"], COORD)
        assert len((await server.list_ambiental("p1", COORD))["items"]) == 1
    _con_db(go)


def test_validaciones_y_permisos():
    async def go(db):
        for malo in (server.AmbientalIn(seccion="informe", nombre="x", avance=150),
                     server.AmbientalIn(seccion="informe", nombre="x", fecha_cierre="nunca"),
                     server.AmbientalIn(seccion="informe", nombre=" "),
                     server.AmbientalIn(seccion="arbolado", nombre="x", poda_aut=-1)):
            with pytest.raises(HTTPException) as e:
                await server.create_ambiental("p1", malo, COORD)
            assert e.value.status_code == 400
        with pytest.raises(HTTPException) as e:  # el Jefe solo lee
            await server.create_ambiental("p1", server.AmbientalIn(seccion="informe", nombre="x"), JEFE)
        assert e.value.status_code in (401, 403)
        with pytest.raises(HTTPException):
            await server.set_ambiental_config("p1", server.AmbientalConfigIn(total_a_restituir=-3), COORD)
    _con_db(go)


def test_modulo_apagado_no_deja_capturar():
    async def go(db):
        with pytest.raises(HTTPException) as e:
            await server.create_ambiental("p1", server.AmbientalIn(seccion="informe", nombre="x"), COORD)
        assert e.value.status_code == 403
    _con_db(go, modulos={"reportes_diarios": {}, "reportes_semanales": {}, "reportes_mensuales": {}})


def test_bloques():
    async def go(db):
        await _sembrar_todo()
        items = await db.ambiental.find({"project_id": "p1"}).to_list(length=100)
        cfg = await db.ambiental_config.find_one({"project_id": "p1"})
        return await server._ambiental_bloques(items, cfg)
    bloques = _con_db(go)
    tipos = [b["type"] for b in bloques]
    assert tipos == ["imagen", "table", "imagen", "imagen", "table", "summary", "fotos_comparativa", "table"]
    assert bloques[0]["heading"].startswith("SEGUIMIENTO AL CUMPLIMIENTO") and "05134/2026" in bloques[0]["caption"]
    assert bloques[1]["rows"] == [["SEDEMA", "SEDEMA/DGEIRA/DEIAR/05134/2026"]]
    assert bloques[2]["heading"] == "INTERVENCIÓN DE ARBOLADO · MIA-E"
    assert bloques[5]["text"] == "Total a restituir: 46 · Total restituidos: 46"
    assert len(bloques[6]["items"]) == 1 and bloques[6]["items"][0]["etiqueta_original"] == "Se plantan 30 ejemplares"
    assert bloques[7]["rows"] == [["Dictamen técnico forestal", "100%"]]


def test_sin_datos_no_imprime():
    async def go(db):
        return await server._ambiental_bloques([], {})
    assert _con_db(go) == []


def _exportar(formato):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await db.projects.insert_one({"id": "p1", "name": "Proyecto Demo", "color_institucional": "#7A1030", "modulos": ACTIVO})
            await _sembrar_todo()
            common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=COORD)
            if formato == "pdf":
                resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                       recuento_sin_datos=None, estructura_solo=False)
            elif formato == "docx":
                resp = await server.export_reports_docx(**common)
            else:
                resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None,
                                                        recuento_sin_datos=None)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def test_export_pdf():
    from pypdf import PdfReader
    txt = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(_exportar("pdf"))).pages)
    for esperado in ("SUPERVISIÓN AMBIENTAL", "AUTORIZACIONES", "RESARCIMIENTO DE ARBOLADO", "Total a restituir: 46",
                     "TABLA RESUMEN GENERAL", "Plantaciones en Parque Cantera"):
        assert esperado in txt, esperado


def test_export_pptx_y_docx():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx")))
    textos = []
    for sl in prs.slides:
        textos += [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        textos += [c.text for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table
                   for r in sh.table.rows for c in r.cells]
    plano = "\n".join(textos)
    for esperado in ("SUPERVISIÓN AMBIENTAL", "SEDEMA/DGEIRA/DEIAR/05134/2026", "Total a restituir: 46", "Dictamen técnico forestal"):
        assert esperado in plano, esperado
    from docx import Document
    d = Document(io.BytesIO(_exportar("docx")))
    plano_d = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    assert "SUPERVISIÓN AMBIENTAL" in plano_d and "Total a restituir: 46" in plano_d
