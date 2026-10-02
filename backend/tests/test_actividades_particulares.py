"""Actividades particulares: catálogo de categorías/subcategorías EDITABLE,
registros por estación e intertramo, y "ACTIVIDADES DESGLOSADAS" en el export
(categorías sin actividad = "Sin actividades ésta semana")."""
import asyncio
import io
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_act_part")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import ROLE_COORD, ROLE_ESPECIALISTA, ROLE_JEFE, ROLE_SUB, ACT_PART_SIN_ACTIVIDAD  # noqa: E402

pytest.importorskip("mongomock_motor")
COORD = {"id": "u1", "role": ROLE_COORD, "name": "C", "email": "c@x.mx", "project_ids": ["p1"]}
SUB = {"id": "u2", "role": ROLE_SUB, "name": "S", "email": "s@x.mx", "project_ids": ["p1"]}
ESP = {"id": "u4", "role": ROLE_ESPECIALISTA, "name": "E", "email": "e@x.mx", "project_ids": ["p1"]}
ESP2 = {"id": "u5", "role": ROLE_ESPECIALISTA, "name": "E2", "email": "e2@x.mx", "project_ids": ["p1"]}
JEFE = {"id": "u3", "role": ROLE_JEFE, "name": "J", "email": "j@x.mx", "project_ids": ["p1"]}
ACTIVO = {"reportes_diarios": {}, "reportes_mensuales": {}, "reportes_semanales": {"actividades_particulares": True}}


async def _base(db):
    await db.projects.insert_one({"id": "p1", "name": "Proyecto Demo", "color_institucional": "#7A1030", "modulos": ACTIVO})
    for i, (nid, nombre, parent, order) in enumerate([("t1", "TRAMO 1", None, 0), ("e1", "ESTACIÓN 1. UNIVERSIDAD", "t1", 0),
                                                     ("e2", "ESTACIÓN 2. CANTERA", "t1", 1)]):
        await db.location_nodes.insert_one({"id": nid, "project_id": "p1", "name": nombre, "parent_id": parent, "depth": 0 if not parent else 1,
                                            "order": order, "is_leaf": bool(parent), "measurement_type": "cadenamiento"})


def _con_db(fn):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await _base(db)
            return await fn(db)
        return asyncio.run(go())
    finally:
        server.db = old


def test_catalogo_editable_sugeridas_sin_duplicar_y_proteccion():
    async def go(db):
        r = await server.cargar_act_part_sugeridas("p1", COORD)
        assert r["agregadas"] == len(server.ACT_PART_SUGERIDAS) and r["items"][0]["nombre"] == "ESTUDIOS Y ANÁLISIS"
        assert (await server.cargar_act_part_sugeridas("p1", COORD))["agregadas"] == 0  # idempotente
        # editable: renombrar, cambiar subcategorías, agregar y reordenar
        c0 = r["items"][0]
        u = await server.update_act_part_cat("p1", c0["id"], server.ActPartCatPatch(nombre="Estudios", subcategorias=[" Suelos ", "Suelos", "", "Topo"]), COORD)
        assert u["nombre"] == "Estudios" and u["subcategorias"] == ["Suelos", "Topo"]
        nueva = await server.create_act_part_cat("p1", server.ActPartCatIn(nombre="Mi categoría", subcategorias=["A"]), SUB)
        assert nueva["orden"] == len(server.ACT_PART_SUGERIDAS) + 1
        with pytest.raises(HTTPException) as e:
            await server.create_act_part_cat("p1", server.ActPartCatIn(nombre="Mi categoría"), COORD)
        assert e.value.status_code == 409
    _con_db(go)


def test_solo_coordinador_y_sub_editan_el_catalogo():
    guard = server.require_role(ROLE_COORD, ROLE_SUB)
    assert asyncio.run(guard(SUB)) == SUB
    with pytest.raises(HTTPException) as e:
        asyncio.run(guard(ESP))
    assert e.value.status_code == 403


def test_registros_crud_renombrar_categoria_y_borrado_protegido():
    async def go(db):
        cat = await server.create_act_part_cat("p1", server.ActPartCatIn(nombre="PROYECTO EJECUTIVO", subcategorias=["General"]), COORD)
        reg = await server.create_act_part("p1", server.ActPartIn(node_id="e1", categoria_id=cat["id"], subcategoria=" General ",
                                                                   texto=" Seguimiento técnico ", fecha="30/09/2026"), ESP)
        assert reg["texto"] == "Seguimiento técnico" and reg["subcategoria"] == "General" and reg["fecha"] == "2026-09-30"
        assert reg["categoria_nombre"] == "PROYECTO EJECUTIVO"
        # editar el nombre de la categoría actualiza lo capturado
        await server.update_act_part_cat("p1", cat["id"], server.ActPartCatPatch(nombre="PROYECTO"), COORD)
        assert (await db.actividades_particulares.find_one({"id": reg["id"]}))["categoria_nombre"] == "PROYECTO"
        # no se puede borrar una categoría con actividades
        with pytest.raises(HTTPException) as e:
            await server.delete_act_part_cat("p1", cat["id"], COORD)
        assert e.value.status_code == 409
        # editable
        u = await server.update_act_part("p1", reg["id"], server.ActPartPatch(texto="Otro texto", node_id="e2"), ESP)
        assert u["texto"] == "Otro texto" and u["node_id"] == "e2"
        # validaciones y permisos
        for malo in (server.ActPartIn(node_id="zzz", categoria_id=cat["id"], texto="x"),
                     server.ActPartIn(node_id="e1", categoria_id="zzz", texto="x"),
                     server.ActPartIn(node_id="e1", categoria_id=cat["id"], texto="  ")):
            with pytest.raises(HTTPException) as e:
                await server.create_act_part("p1", malo, COORD)
            assert e.value.status_code == 400
        with pytest.raises(HTTPException):
            await server.create_act_part("p1", server.ActPartIn(node_id="e1", categoria_id=cat["id"], texto="x"), JEFE)
        with pytest.raises(HTTPException) as e:  # otro especialista no borra lo ajeno
            await server.delete_act_part("p1", reg["id"], ESP2)
        assert e.value.status_code == 403
        await server.delete_act_part("p1", reg["id"], SUB)
        await server.delete_act_part_cat("p1", cat["id"], COORD)  # ya sin actividades: se puede
    _con_db(go)


def test_bloques_por_estacion_con_todas_las_categorias():
    cat = [{"id": "c1", "nombre": "ESTUDIOS Y ANÁLISIS", "subcategorias": ["Mecánica de Suelos", "Topografía"], "orden": 1},
           {"id": "c2", "nombre": "OBRAS INDUCIDAS", "subcategorias": [], "orden": 2}]
    nodos = [{"id": "t1", "name": "T", "parent_id": None, "order": 0}, {"id": "e1", "name": "ESTACIÓN 1. UNIVERSIDAD", "parent_id": "t1", "order": 0},
             {"id": "e2", "name": "ESTACIÓN 2. CANTERA", "parent_id": "t1", "order": 1}]
    regs = [
        {"node_id": "e2", "categoria_id": "c2", "texto": "Reunión con CIS", "fecha": "2026-09-22", "created_at": 1},
        {"node_id": "e1", "categoria_id": "c1", "subcategoria": "Topografía", "texto": "Asistencia a reunión", "fecha": "2026-09-22", "created_at": 2},
        {"node_id": "e1", "categoria_id": "c1", "subcategoria": "Mecánica de Suelos", "texto": "Revisión del EMS", "fecha": "2026-09-23", "created_at": 3},
        {"node_id": "e1", "categoria_id": "c1", "subcategoria": "Mecánica de Suelos", "texto": "Se emite informe", "fecha": "2026-09-24", "created_at": 4},
        {"node_id": "e1", "categoria_id": "borrada", "categoria_nombre": "VIEJA", "texto": "Dato antiguo", "fecha": "2026-09-24", "created_at": 5},
    ]
    b = server._actividades_particulares_bloques(regs, cat, nodos)
    assert [x["heading"] for x in b] == ["ESTACIÓN 1. UNIVERSIDAD · ACTIVIDADES DESGLOSADAS", "ESTACIÓN 2. CANTERA · ACTIVIDADES DESGLOSADAS"]
    assert b[0]["rows"] == [
        ["ESTUDIOS Y ANÁLISIS", "Mecánica de Suelos", "• Revisión del EMS"],
        ["", "", "• Se emite informe"],
        ["", "Topografía", "• Asistencia a reunión"],
        ["OBRAS INDUCIDAS", "", ACT_PART_SIN_ACTIVIDAD],
        ["VIEJA", "", "• Dato antiguo"],
    ]
    assert b[1]["rows"][0] == ["ESTUDIOS Y ANÁLISIS", "", ACT_PART_SIN_ACTIVIDAD] and b[1]["rows"][1] == ["OBRAS INDUCIDAS", "", "• Reunión con CIS"]
    assert server._actividades_particulares_bloques([], cat, nodos) == []


def _exportar(formato):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await _base(db)
            await server.cargar_act_part_sugeridas("p1", COORD)
            cat = (await server.list_act_part_cat("p1", COORD))[0]
            await server.create_act_part("p1", server.ActPartIn(node_id="e1", categoria_id=cat["id"], subcategoria="Topografía",
                                                                 texto="Asistencia a reunión de seguimiento"), COORD)
            common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=COORD)
            if formato == "pdf":
                resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                       recuento_sin_datos=None, estructura_solo=False)
            elif formato == "docx":
                resp = await server.export_reports_docx(**common)
            else:
                resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None,
                                                        recuento_sin_datos=None)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def test_export_los_tres_formatos():
    from pypdf import PdfReader
    pdf = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(_exportar("pdf"))).pages)
    for esperado in ("ACTIVIDADES PARTICULARES", "ACTIVIDADES DESGLOSADAS", "Asistencia a reunión de seguimiento", ACT_PART_SIN_ACTIVIDAD, "PROYECTO EJECUTIVO"):
        assert esperado in pdf, esperado
    from pptx import Presentation
    textos = []
    for sl in Presentation(io.BytesIO(_exportar("pptx"))).slides:
        textos += [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        textos += [c.text for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table for r in sh.table.rows for c in r.cells]
    plano = "\n".join(textos)
    assert "ESTACIÓN 1. UNIVERSIDAD · ACTIVIDADES DESGLOSADAS" in plano and "Asistencia a reunión de seguimiento" in plano
    from docx import Document
    d = Document(io.BytesIO(_exportar("docx")))
    plano_d = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    assert "ACTIVIDADES DESGLOSADAS" in plano_d and "Asistencia a reunión de seguimiento" in plano_d
