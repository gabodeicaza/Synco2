"""Indicadores de seguridad e higiene del reporte semanal (server.py):
cumplimiento en campo, documental (mensual y por semana) y PIPC por estación.
Tests puros y de punta a punta del export en PDF/PPTX/DOCX contra una Mongo
simulada (`mongomock-motor`, se omite si no está instalado)."""
import asyncio
import io
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_indicadores_seguridad")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import (  # noqa: E402
    IndicadorIn, _indicador_normalizar, _indicadores_seguridad_graficas, _pipc_bloques, _seg_campo_resumen,
    _seg_dias_sin_accidentes, _seg_doc_mensual, _seg_doc_semanal,
)

CORTE = date(2026, 6, 27)


def _png_ok(png: bytes):
    from PIL import Image
    assert Image.open(io.BytesIO(png)).size == (1800, 800)


# ── Normalización / validación ─────────────────────────────────────────────
def test_normalizar_cada_tipo():
    assert _indicador_normalizar(IndicadorIn(tipo="campo", fecha="2026-06-27", porcentaje=42)) == (
        "campo|2026-06-27", {"fecha": "2026-06-27", "porcentaje": 42.0, "dias_sin_accidentes": None})
    assert _indicador_normalizar(IndicadorIn(tipo="doc_mensual", mes="2026-05", porcentaje=67.37))[0] == "doc_mensual|2026-05"
    # cualquier día de la semana se normaliza a su lunes (27-jun-2026 es sábado → lunes 22)
    clave, campos = _indicador_normalizar(IndicadorIn(tipo="doc_semanal", semana="2026-06-27", categoria="seguridad", porcentaje=65))
    assert clave == "doc_semanal|2026-06-22|seguridad" and campos["semana"] == "2026-06-22"
    clave, campos = _indicador_normalizar(IndicadorIn(tipo="pipc", etiqueta="e1", puntos=[51, 0, 0, 0], orden=1))
    assert clave == "pipc|E1" and campos["puntos"] == [51.0, 0.0, 0.0, 0.0]
    assert _indicador_normalizar(IndicadorIn(tipo="nota_pipc", texto="  hola  ")) == ("nota_pipc", {"texto": "hola"})


@pytest.mark.parametrize("body", [
    dict(tipo="campo", fecha="2026-06-27", porcentaje=101),
    dict(tipo="campo", fecha="2026-06-27", porcentaje=None),
    dict(tipo="campo", fecha="27/13/2026", porcentaje=10),
    dict(tipo="campo", fecha="2026-06-27", porcentaje=10, dias_sin_accidentes=-1),
    dict(tipo="doc_mensual", mes="2026-13", porcentaje=10),
    dict(tipo="doc_mensual", mes="mayo", porcentaje=10),
    dict(tipo="doc_semanal", semana="2026-06-22", porcentaje=10),  # falta categoría
    dict(tipo="pipc", etiqueta="E1", puntos=[1, 2, 3]),  # faltan pasos
    dict(tipo="pipc", etiqueta="E1", puntos=[71, 0, 0, 0]),  # elaboración > 70
    dict(tipo="pipc", etiqueta="E1", puntos=[10, 20, 10, 5]),  # pasos 2-4 > 30
    dict(tipo="pipc", etiqueta="E1", puntos=[10, -1, 0, 0]),
    dict(tipo="pipc", etiqueta="  ", puntos=[0, 0, 0, 0]),
])
def test_normalizar_rechaza_valores_invalidos(body):
    with pytest.raises(HTTPException) as e:
        _indicador_normalizar(IndicadorIn(**body))
    assert e.value.status_code == 400


# ── Cálculos ───────────────────────────────────────────────────────────────
def test_campo_actual_y_anterior_al_corte():
    regs = [{"tipo": "campo", "fecha": "2026-06-13", "porcentaje": 39.2}, {"tipo": "campo", "fecha": "2026-06-27", "porcentaje": 42.0},
            {"tipo": "campo", "fecha": "2026-07-04", "porcentaje": 90.0}]  # posterior al corte: no cuenta
    r = _seg_campo_resumen(regs, CORTE)
    assert (r["actual"], r["anterior"], r["fecha"]) == (42.0, 39.2, "2026-06-27")
    assert _seg_campo_resumen(regs[:1], CORTE)["anterior"] is None
    assert _seg_campo_resumen([], CORTE) is None


def _acc(fecha, tipo):
    return {"fecha": fecha, "tipo_registro": tipo, "descripcion": "x"}


def test_dias_sin_accidentes_solo_reinician_incapacitante_y_grave():
    reportes = [_acc("2026-06-01", "incapacitante"), _acc("2026-06-24", "accidente_leve"),
                _acc("2026-06-25", "accidente_leve"), _acc("2026-06-26", "incidente"), _acc("2026-06-28", "grave")]
    assert _seg_dias_sin_accidentes(reportes, CORTE) == 26  # 1-jun → 27-jun; el leve y lo posterior no reinician
    assert _seg_dias_sin_accidentes(reportes, CORTE, manual=9) == 9  # lo manual manda
    # sin accidentes que reinicien: desde el primer registro
    assert _seg_dias_sin_accidentes([_acc("2026-06-20", "sin_accidentes")], CORTE) == 7
    assert _seg_dias_sin_accidentes([], CORTE) is None


def test_documental_mensual_y_semanal():
    regs = [{"tipo": "doc_mensual", "mes": "2026-04", "porcentaje": 42}, {"tipo": "doc_mensual", "mes": "2026-05", "porcentaje": 67.37},
            {"tipo": "doc_mensual", "mes": "2026-07", "porcentaje": 99}]
    m = _seg_doc_mensual(regs, CORTE)
    assert (m["actual"]["mes"], m["anterior"]["mes"]) == ("2026-05", "2026-04")
    sem = [{"tipo": "doc_semanal", "semana": s, "categoria": c, "porcentaje": p}
           for s, vals in (("2026-06-08", (10, 10, 10)), ("2026-06-15", (47, 45, 35)), ("2026-06-22", (65, 63, 74)))
           for c, p in zip(("seguridad", "proteccion_civil", "higiene_salud"), vals)]
    r = _seg_doc_semanal(sem, CORTE)
    assert [x["label"] for x in r["semanas"]] == ["SEMANA 25", "SEMANA 26"]  # solo las 2 últimas
    assert r["categorias"][0] == {"clave": "seguridad", "label": "Seguridad", "valores": [47, 65]}
    assert _seg_doc_semanal([], CORTE) is None


def test_pipc_tabla_totales_y_notas():
    regs = [{"tipo": "pipc", "etiqueta": f"E{i}", "orden": i, "puntos": [51, 0, 0, 0]} for i in range(1, 9)]
    regs.append({"tipo": "pipc", "etiqueta": "E9", "orden": 9, "puntos": [70, 10, 10, 10]})
    regs.append({"tipo": "nota_pipc", "texto": "Se reciben los estudios\n\n  Se solicita precisar los intertramos "})
    tabla, nota, notas = _pipc_bloques(regs)
    assert tabla["headers"][:3] == ["N°", "PASOS - ETAPAS", "CONCEPTO"] and tabla["headers"][3:] == [f"E{i}" for i in range(1, 10)]
    assert tabla["rows"][0][1] == "ELABORACIÓN" and tabla["rows"][0][3] == "51" and tabla["rows"][0][-1] == "70"
    assert tabla["rows"][-1][3] == "51%" and tabla["rows"][-1][-1] == "100%"
    assert "Promedio de las estaciones: 56.44%" in tabla["rows"][-1][2]  # (8*51+100)/9
    assert all(len(r) == len(tabla["headers"]) for r in tabla["rows"]) and abs(sum(tabla["col_widths"]) - 1) < 1e-9
    assert "70 puntos" in nota["items"][0]["body"]
    assert [i["body"] for i in notas["items"]] == ["Se reciben los estudios", "Se solicita precisar los intertramos"]
    assert _pipc_bloques([]) == []


def test_graficas_y_bloques_se_omiten_sin_datos_y_son_png_validos():
    assert _indicadores_seguridad_graficas([], [], CORTE) == ([], [], [])
    regs = [{"tipo": "campo", "fecha": "2026-06-27", "porcentaje": 42.0},
            {"tipo": "doc_mensual", "mes": "2026-05", "porcentaje": 67.37},
            {"tipo": "doc_semanal", "semana": "2026-06-22", "categoria": "seguridad", "porcentaje": 65}]
    bloques, graficas, despues = _indicadores_seguridad_graficas(regs, [_acc("2026-06-18", "incapacitante")], CORTE)
    assert bloques == [{"type": "summary", "text": "DÍAS SIN ACCIDENTES: 9"}] and despues == []
    assert [g[0] for g in graficas] == ["CUMPLIMIENTO EN CAMPO", "CUMPLIMIENTO DOCUMENTAL", "PORCENTAJE DE CUMPLIMIENTO POR ÁREA"]
    for _, _, png in graficas:
        _png_ok(png)
    assert "Porcentaje actual 42%" in graficas[0][1]


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
SEG = {"id": "u2", "role": server.ROLE_ESPECIALISTA, "name": "Seguridad", "email": "s@x.mx", "project_ids": ["p1"]}
OTRO = {"id": "u3", "role": server.ROLE_ESPECIALISTA, "name": "Otro", "email": "o@x.mx", "project_ids": ["p1"]}
JEFE = {"id": "u4", "role": server.ROLE_JEFE, "name": "Jefe", "email": "j@x.mx"}
MODULOS = {"indicadores_seguridad": True, "pendientes_generales": True}
HOY = datetime.now(server.MX_TZ).date()


def _con_db(coro_fn):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        return asyncio.run(coro_fn(db))
    finally:
        server.db = old


async def _sembrar(db):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": {}, "reportes_semanales": MODULOS, "reportes_mensuales": {}},
    })
    put = server.set_seguridad_indicador
    hoy, mes = HOY.isoformat(), f"{HOY.year:04d}-{HOY.month:02d}"
    previo = (HOY.replace(day=1) - timedelta(days=1))
    await put("p1", IndicadorIn(tipo="campo", fecha=(HOY - timedelta(days=14)).isoformat(), porcentaje=39.2), SEG)
    await put("p1", IndicadorIn(tipo="campo", fecha=hoy, porcentaje=42), SEG)
    await put("p1", IndicadorIn(tipo="doc_mensual", mes=f"{previo.year:04d}-{previo.month:02d}", porcentaje=42), SEG)
    await put("p1", IndicadorIn(tipo="doc_mensual", mes=mes, porcentaje=67.37), SEG)
    for sem, vals in (((HOY - timedelta(days=7)).isoformat(), (47, 45, 35)), (hoy, (65, 63, 74))):
        for cat, v in zip(("seguridad", "proteccion_civil", "higiene_salud"), vals):
            await put("p1", IndicadorIn(tipo="doc_semanal", semana=sem, categoria=cat, porcentaje=v), SEG)
    for i in range(1, 4):
        await put("p1", IndicadorIn(tipo="pipc", etiqueta=f"E{i}", orden=i, puntos=[51, 0, 0, 0]), SEG)
    await put("p1", IndicadorIn(tipo="nota_pipc", texto="La supervisión externa recibe los Estudios de Riesgo de las ocho estaciones"), SEG)
    await server.create_pendiente_general("p1", server.PendienteIn(
        contrato="seguridad", titulo="Actualización del SSySO", descripcion="conforme a la NOM-031-STPS-2011"), SEG)
    await db.seguridad_reportes.insert_one({"id": "a1", "project_id": "p1", "fecha": (HOY - timedelta(days=9)).isoformat(),
                                            "tipo_registro": "incapacitante", "descripcion": "Caída", "created_at": datetime.now(server.timezone.utc)})


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_semanal_imprime_los_indicadores(formato):
    from test_capa_financiera import _texto

    async def go(db):
        await _sembrar(db)
        common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER)
        if formato == "pdf":
            resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                   recuento_sin_datos=None, estructura_solo=False)
        elif formato == "pptx":
            resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
        else:
            resp = await server.export_reports_docx(**common)
        return b"".join([c async for c in resp.body_iterator])
    data = _con_db(go)
    texto = _texto(formato, data)
    for esperado in (
        "CUMPLIMIENTO SEGURIDAD E HIGIENE", "DÍAS SIN ACCIDENTES: 9", "CUMPLIMIENTO EN CAMPO", "Porcentaje actual 42%",
        "Porcentaje anterior 39.2%", "CUMPLIMIENTO DOCUMENTAL", "67.37%", "PORCENTAJE DE CUMPLIMIENTO POR ÁREA",
        "CUMPLIMIENTO DEL PIPC MODALIDAD ESTUDIO DE RIESGO OBRA Y DEMOLICIÓN (EROD)", "ELABORACIÓN", "CAPTURA DIGITAL",
        "OBTENCIÓN DE FOLIO", "CUMPLIMIENTO TOTAL", "51%", "Notas de avance", "recibe los Estudios de Riesgo",
        "PENDIENTES ADMINISTRATIVOS (SEGURIDAD E HIGIENE)", "Actualización del SSySO",
    ):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"
    if formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 3  # medidor + dona + barras
        assert not [1 for sl in prs.slides for sh in sl.shapes if sh.top + sh.height > prs.slide_height + 5_000]


def test_upsert_permisos_y_modulo_apagado():
    async def go(db):
        await _sembrar(db)
        n = len(await server.list_seguridad_indicadores("p1", SEG))
        # mismo corte → reemplaza, no duplica
        await server.set_seguridad_indicador("p1", IndicadorIn(tipo="campo", fecha=HOY.isoformat(), porcentaje=50), SEG)
        lista = await server.list_seguridad_indicadores("p1", SEG)
        assert len(lista) == n and next(r for r in lista if r["tipo"] == "campo" and r["fecha"] == HOY.isoformat())["porcentaje"] == 50
        with pytest.raises(HTTPException) as e:
            await server.set_seguridad_indicador("p1", IndicadorIn(tipo="campo", fecha=HOY.isoformat(), porcentaje=1), JEFE)
        assert e.value.status_code == 403
        rid = lista[0]["id"]
        with pytest.raises(HTTPException) as e2:
            await server.delete_seguridad_indicador("p1", rid, OTRO)
        assert e2.value.status_code == 403
        assert (await server.delete_seguridad_indicador("p1", rid, USER))["ok"]
        await db.projects.update_one({"id": "p1"}, {"$set": {"modulos.reportes_semanales": {}}})
        with pytest.raises(HTTPException) as e3:
            await server.set_seguridad_indicador("p1", IndicadorIn(tipo="campo", fecha=HOY.isoformat(), porcentaje=1), SEG)
        assert e3.value.status_code == 403
    _con_db(go)
