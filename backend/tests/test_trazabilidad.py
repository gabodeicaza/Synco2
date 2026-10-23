"""Trazabilidad (server.py): árbol categorías → grupos → documentos que reemplaza a
los "Archivos de consulta". Importación desde el Excel de trazabilidad (mezcla, no
reemplazo), edición de características, movimiento de nodos y archivos adjuntos.
Contra una Mongo simulada (`mongomock-motor`)."""
import asyncio
import io
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_trazabilidad")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402

pytest.importorskip("mongomock_motor")

COORD = {"id": "c1", "role": server.ROLE_COORD, "name": "Coord"}
CLAVE = "DGOT-L4CB-ALF-LEIT-AP-TR1-TOP-PLN-E1-1201-R00"


def _con_db(coro_fn):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old_db, old_up = server.db, server.UPLOADS_DIR
    server.db = db
    with tempfile.TemporaryDirectory() as tmp:
        server.UPLOADS_DIR = Path(tmp)
        try:
            return asyncio.run(coro_fn(db))
        finally:
            server.db, server.UPLOADS_DIR = old_db, old_up


def _subida(nombre, contenido):
    async def _leer():
        return contenido
    return SimpleNamespace(filename=nombre, read=_leer)


def _excel(oficio="DIR-CL4-150-2026", extra_plano=False) -> bytes:
    """Hoja mínima con la forma de 'Trazabilidad PE': encabezado en filas 14-15,
    A101 > A1011 > ESTACIONES > Estación E1 > planos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trazabilidad PE (2)"
    ws["B14"], ws["C14"], ws["D14"] = "CLAVE DEL PLANO", "NOMBRE DEL PLANO", "DISCIPLINA"
    ws["E14"], ws["H14"], ws["P14"], ws["N14"] = "ESCRITO ALFA A DGOT\nREV. A", "ENVIO DGOT A DIRAC", "VALIDACIÓN TÉCNICA POR DIRAC", "ULTIMA REV. PDF"
    ws["AA14"] = "OBSERVACIONES"
    for col, txt in (("E", "REV"), ("F", "OFICIO"), ("G", "FECHA"), ("P", "REV"), ("Q", "OFICIO"), ("R", "FECHA")):
        ws[f"{col}15"] = txt
    filas = [
        ("A101", "ESTUDIOS Y ANÁLISIS", 0), ("A1011", "Estudios Topográficos.", 1),
        (None, "ESTACIONES", 2), (None, "Estación E1 - Universidad", 2),
    ]
    r = 16
    for b, c, nivel in filas:
        ws.cell(r, 2, b)
        ws.cell(r, 3, c)
        ws.row_dimensions[r].outlineLevel = nivel
        r += 1
    ws.cell(r, 2, CLAVE)
    ws.cell(r, 3, "ALTIMETRÍA")
    ws.cell(r, 4, "ALTIMETRÍA")
    ws.cell(r, 5, "A")
    ws.cell(r, 6, "APYCSA-499-2026")
    ws.cell(r, 7, datetime(2026, 8, 25))
    ws.cell(r, 14, '=HYPERLINK(".\\01\\"& B20&".PDF","R00")')
    ws.cell(r, 16, 0)
    ws.cell(r, 17, oficio)
    ws.cell(r, 18, datetime(2026, 4, 29))
    ws.cell(r, 27, "Pendiente firma")
    if extra_plano:
        ws.cell(r + 1, 2, CLAVE.replace("1201", "1301"))
        ws.cell(r + 1, 3, "CURVAS DE NIVEL")
    ws.cell(r + 3, 4, "TOTAL")
    ws.cell(r + 5, 3, "RESUMEN")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _importar(excel, **kw):
    return await server.traz_importar("p1", _subida("traz.xlsx", excel), user=COORD, **{"hoja": None, "solo_vacios": False, **kw})


async def _sembrar(db):
    await db.projects.insert_one({"id": "p1", "name": "Demo"})


def _por_codigo(nodes, code):
    return next(n for n in nodes if n["code"] == code)


def test_importa_el_arbol_con_jerarquia_y_caracteristicas():
    async def go(db):
        await _sembrar(db)
        res = await _importar(_excel())
        return res, (await server.traz_listar("p1", COORD))["nodes"]
    res, nodes = _con_db(go)
    assert res["creados"] == 5 and res["actualizados"] == 0
    a101, a1011 = _por_codigo(nodes, "A101"), _por_codigo(nodes, "A1011")
    assert a101["parent_id"] is None and a1011["parent_id"] == a101["id"]
    estaciones = next(n for n in nodes if n["name"] == "ESTACIONES")
    e1 = next(n for n in nodes if n["name"].startswith("Estación E1"))
    assert estaciones["parent_id"] == a1011["id"] and e1["parent_id"] == estaciones["id"]
    doc = _por_codigo(nodes, CLAVE)
    assert doc["kind"] == "documento" and doc["parent_id"] == e1["id"]
    f = doc["fields"]
    assert f["alfa_rev"] == "A" and f["alfa_fecha"] == "2026-08-25" and f["pdf"] == "R00"
    assert f["val_oficio"] == "DIR-CL4-150-2026" and f["val_fecha"] == "2026-04-29" and f["obs"] == "Pendiente firma"


def test_reimportar_actualiza_sin_duplicar_ni_perder_lo_capturado():
    async def go(db):
        await _sembrar(db)
        await _importar(_excel())
        nodes = (await server.traz_listar("p1", COORD))["nodes"]
        doc = _por_codigo(nodes, CLAVE)
        await server.traz_editar("p1", doc["id"], server.TrazNodoPatch(fields={"estatus": "Firmado"}), COORD)
        res = await _importar(_excel(oficio="DIR-CL4-999-2026", extra_plano=True))
        return res, (await server.traz_listar("p1", COORD))["nodes"]
    res, nodes = _con_db(go)
    assert res["creados"] == 1 and res["actualizados"] == 1
    assert len([n for n in nodes if n["kind"] == "documento"]) == 2
    f = _por_codigo(nodes, CLAVE)["fields"]
    assert f["val_oficio"] == "DIR-CL4-999-2026" and f["estatus"] == "Firmado"


def test_editar_caracteristicas_fechas_y_valores_vacios():
    async def go(db):
        await _sembrar(db)
        await _importar(_excel())
        doc = _por_codigo((await server.traz_listar("p1", COORD))["nodes"], CLAVE)
        r = await server.traz_editar("p1", doc["id"], server.TrazNodoPatch(
            fields={"dgot_fecha": "02/09/2026", "alfa_rev": "", "disciplina": "TOPO"},
            extra=[{"label": "Responsable", "value": "Ing. Pérez"}, {"label": " ", "value": "x"}]), COORD)
        with pytest.raises(HTTPException):
            await server.traz_editar("p1", doc["id"], server.TrazNodoPatch(fields={"dgot_fecha": "31/02/2026"}), COORD)
        with pytest.raises(HTTPException):
            await server.traz_editar("p1", doc["id"], server.TrazNodoPatch(fields={"inventada": "x"}), COORD)
        return r
    r = _con_db(go)
    assert r["fields"]["dgot_fecha"] == "2026-09-02" and "alfa_rev" not in r["fields"]
    assert r["fields"]["disciplina"] == "TOPO"
    assert r["extra"] == [{"label": "Responsable", "value": "Ing. Pérez"}]


def test_crear_mover_y_borrar_nodos():
    async def go(db):
        await _sembrar(db)
        cat = await server.traz_crear("p1", server.TrazNodoIn(kind="categoria", code="A401", name="Nueva"), COORD)
        doc = await server.traz_crear("p1", server.TrazNodoIn(parent_id=cat["id"], code="PLANO-1", name="Uno"), COORD)
        otra = await server.traz_crear("p1", server.TrazNodoIn(kind="categoria", code="A402", name="Otra"), COORD)
        with pytest.raises(HTTPException):  # no se puede meter una categoría dentro de su hijo
            await server.traz_editar("p1", cat["id"], server.TrazNodoPatch(parent_id=doc["id"]), COORD)
        movido = await server.traz_editar("p1", doc["id"], server.TrazNodoPatch(parent_id=otra["id"]), COORD)
        raiz = await server.traz_editar("p1", otra["id"], server.TrazNodoPatch(parent_id=None), COORD)
        borrado = await server.traz_borrar("p1", otra["id"], COORD)
        quedan = (await server.traz_listar("p1", COORD))["nodes"]
        return movido, raiz, borrado, quedan, cat
    movido, raiz, borrado, quedan, cat = _con_db(go)
    assert movido["parent_id"] and raiz["parent_id"] is None
    assert borrado["borrados"] == 2 and [n["id"] for n in quedan] == [cat["id"]]


def test_archivos_adjuntar_reemplazar_descargar_y_acomodo_automatico():
    async def go(db):
        await _sembrar(db)
        await _importar(_excel(extra_plano=True))
        doc = _por_codigo((await server.traz_listar("p1", COORD))["nodes"], CLAVE)
        r1 = await server.traz_subir_archivo("p1", doc["id"], _subida("rev.pdf", b"%PDF-1"), COORD)
        r2 = await server.traz_subir_archivo("p1", doc["id"], _subida("rev.pdf", b"%PDF-2"), COORD)
        with pytest.raises(HTTPException):
            await server.traz_subir_archivo("p1", doc["id"], _subida("virus.exe", b"MZ"), COORD)
        auto = await server.traz_subir_auto("p1", _subida(CLAVE + ".DWG", b"AC1032"), COORD)
        sin = await server.traz_subir_auto("p1", _subida("otra-cosa.pdf", b"x"), COORD)
        nodo = _por_codigo((await server.traz_listar("p1", COORD))["nodes"], CLAVE)
        resp = await server.traz_descargar("p1", nodo["files"][0]["file_id"], COORD)
        contenido = Path(resp.path).read_bytes()
        borrado = await server.traz_borrar_archivo("p1", nodo["files"][0]["file_id"], COORD)
        return r1, r2, auto, sin, nodo, contenido, borrado
    r1, r2, auto, sin, nodo, contenido, borrado = _con_db(go)
    assert r1["reemplazado"] is False and r2["reemplazado"] is True
    assert auto["matched"] is True and auto["node_code"] == CLAVE
    assert sin["matched"] is False
    assert sorted(f["ext"] for f in nodo["files"]) == [".dwg", ".pdf"]
    assert contenido == b"%PDF-2"
    assert [f["ext"] for f in borrado["node"]["files"]] == [".dwg"]
    assert "path" not in nodo["files"][0]
