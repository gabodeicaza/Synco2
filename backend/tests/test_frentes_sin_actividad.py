"""Frentes (Estación/Intertramo) sin reportes en el período, en el PPTX: se
imprimen igual con "Sin actividades", en su lugar del orden del árbol (el reporte
semanal de referencia muestra siempre todos los frentes)."""
import asyncio
import io
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_frentes_sin_actividad")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402

import server  # noqa: E402

pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
NODOS = [
    ("r1", "TRAMO I", None, 0, False, None),
    ("e1", "ESTACIÓN 1. UNIVERSIDAD", "r1", 1, True, None),
    ("e2", "ESTACIÓN 2. CANTERA", "r1", 1, True, None),
    ("i12", "INTERTRAMO E1-E2", "r1", 1, False, "intertramo"),
    ("s3", "Poste S3", "i12", 2, True, None),
    ("e3", "ESTACIÓN 3. PERISUR", "r1", 1, True, None),
]


async def _sembrar(db, con_reportes: bool):
    await db.projects.insert_one({"id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
                                  "modulos": {"reportes_diarios": {}, "reportes_semanales": {}, "reportes_mensuales": {}}})
    for i, (nid, name, parent, depth, leaf, tipo) in enumerate(NODOS):
        doc = {"id": nid, "project_id": "p1", "name": name, "parent_id": parent, "depth": depth, "order": i,
               "is_leaf": leaf, "measurement_type": "cadenamiento"}
        if tipo:
            doc["tipo_nodo"] = tipo
        await db.location_nodes.insert_one(doc)
    await db.areas.insert_one({"id": "a1", "project_id": "p1", "name": "OBRA CIVIL"})
    if not con_reportes:
        return
    paths = {"e1": ["TRAMO I", "ESTACIÓN 1. UNIVERSIDAD"], "s3": ["TRAMO I", "INTERTRAMO E1-E2", "Poste S3"]}
    for n, (nid, act) in enumerate((("e1", "Armado de columna leitner"), ("s3", "reconstrucción de muro"))):
        await db.reports.insert_one({
            "id": f"r{n}", "project_id": "p1", "node_id": nid, "area_id": "a1", "created_at": datetime.now(timezone.utc),
            "severidad": "importante", "actividades": act, "images": [], "personnel": ["3 Ayudante"], "equipment": [],
            "captured_by": "u1", "captured_by_name": "Ing", "node_path_names": paths[nid], "title": act,
            "measurement_value": {"cadenamiento": "1+000"},
        })


def _slides(con_reportes: bool):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await _sembrar(db, con_reportes)
            resp = await server.export_reports_pptx(
                pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER,
                recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
            return b"".join([c async for c in resp.body_iterator])
        data = asyncio.run(go())
    finally:
        server.db = old
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    return ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


def _idx(slides, *fragmentos):
    return [i for i, t in enumerate(slides) if all(f in t for f in fragmentos)]


def test_frentes_sin_reportes_salen_con_sin_actividad_en_su_lugar_del_arbol():
    slides = _slides(con_reportes=True)
    (i1,) = _idx(slides, "ESTACIÓN 1.", "UNIVERSIDAD", "ACCESO AL PREDIO")
    (i2,) = _idx(slides, "ESTACIÓN 2.", "CANTERA", "ACCESO AL PREDIO", "Sin actividades")
    (i12,) = _idx(slides, "INTERTRAMO E1-E2", "FUERZA DE TRABAJO", "ACTIVIDAD EN POSTES")
    (i3,) = _idx(slides, "ESTACIÓN 3.", "PERISUR", "ACCESO AL PREDIO", "Sin actividades")
    # orden del árbol: E1, E2 (sin reportes), intertramo E1-E2, E3 (sin reportes)
    assert i1 < i2 < i12 < i3
    # los frentes con reportes NO dicen "Sin actividades"
    assert "Sin actividades" not in slides[i1] and "Sin actividades" not in slides[i12]
    # cada frente sin reportes se imprime UNA sola vez
    assert len(_idx(slides, "ESTACIÓN 2.", "CANTERA", "ACCESO AL PREDIO")) == 1
    assert len(_idx(slides, "ESTACIÓN 3.", "PERISUR", "ACCESO AL PREDIO")) == 1


def test_el_resumen_inicial_tambien_lista_los_frentes_sin_actividad():
    slides = _slides(con_reportes=True)
    resumen = next(t for t in slides if t.startswith("ACTIVIDADES GENERALES"))
    assert "ESTACIÓN 2. CANTERA\nSin actividad" in resumen.replace("\r", "")
    assert "Poste S3 reconstrucción de muro" in resumen


def test_sin_ningun_reporte_no_se_imprimen_diapositivas_de_sin_actividad():
    slides = _slides(con_reportes=False)
    assert not _idx(slides, "Sin actividades", "ACCESO AL PREDIO")
    assert any("Sin reportes en el período" in t for t in slides)


def test_recursos_generales_de_la_obra_con_subtitulos_y_sin_el_titulo_anterior():
    slides = _slides(con_reportes=True)
    (i,) = _idx(slides, "RECURSOS GENERALES DE LA OBRA", "FRENTE")
    assert "FUERZA DE TRABAJO" in slides[i] and "MAQUINARIA Y EQUIPO" in slides[i]
    assert not any("RESUMEN POR FRENTE" in t for t in slides)
    indice = slides[1]
    assert "Recursos generales de la obra" in indice and "Resumen por frente" not in indice


def test_no_se_imprime_portadilla_por_nodo():
    """Una diapositiva cuyo único texto es la ruta del nodo ("TRAMO I › ESTACIÓN 1…") es la
    portadilla por nodo, que se quitó a pedido del usuario. La portadilla de Área ("OBRA CIVIL") sí se conserva."""
    slides = _slides(con_reportes=True)
    rutas = {"TRAMO I › ESTACIÓN 1. UNIVERSIDAD", "TRAMO I › INTERTRAMO E1-E2 › Poste S3"}
    assert not [t for t in slides if t.strip() in rutas]
    assert any(t.strip() == "OBRA CIVIL" for t in slides)
    indice = slides[1]
    assert "TRAMO I › ESTACIÓN 1. UNIVERSIDAD · Obra Civil" in indice  # el índice sigue apuntando al bloque del nodo
