"""Capa financiera del reporte semanal (server.py): avance de obra por partidas
(calculado del motor EVM), estimaciones y pendientes generales. Tests puros y
de punta a punta del export en PDF/PPTX/DOCX contra Mongo simulada."""
import asyncio
import io
import os
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_capa_financiera")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

import server  # noqa: E402
from server import (  # noqa: E402
    _avance_obra_bloques, _avance_obra_graficas, _avance_obra_resumen, _estimacion_valor,
    _estimaciones_bloques, _grafica_lineas_png, _pendientes_bloques,
)

HOY = datetime.now(server.MX_TZ).date()
INI, FIN = HOY - timedelta(days=6), HOY


def _concepto(cid, partida, importe, ini, fin):
    return {"id": cid, "project_id": "p1", "partida_id": partida, "activo": True,
            "meta_cantidad": "1", "importe_total_editable": str(importe),
            "fecha_inicio": ini.isoformat(), "fecha_fin": fin.isoformat(),
            "cantidad": "1", "precio_unitario": "0"}


def _datos():
    partidas = [
        {"id": "pa", "clave": "1", "nombre": "ESTUDIOS", "orden": 1, "activo": True},
        {"id": "pb", "clave": "2", "nombre": "OBRA CIVIL", "orden": 2, "activo": True},
        {"id": "pb1", "clave": "2.1", "nombre": "Cimentación", "orden": 1, "activo": True, "partida_padre_id": "pb"},
    ]
    ini = HOY - timedelta(days=99)
    conceptos = [_concepto("c1", "pa", 1000, ini, HOY + timedelta(days=100)),
                 _concepto("c2", "pb1", 3000, ini, HOY + timedelta(days=100))]
    registros = [
        {"concepto_id": "c1", "fecha": (HOY - timedelta(days=30)).isoformat(), "importe_ganado_congelado": "200"},
        {"concepto_id": "c2", "fecha": (HOY - timedelta(days=2)).isoformat(), "importe_ganado_congelado": "100"},
        {"concepto_id": "c2", "fecha": (HOY - timedelta(days=20)).isoformat(), "importe_ganado_congelado": "300"},
        {"concepto_id": "c2", "fecha": (HOY + timedelta(days=5)).isoformat(), "importe_ganado_congelado": "999"},  # futuro
    ]
    gastos = [{"fecha": (HOY - timedelta(days=10)).isoformat(), "monto": "500"}]
    return conceptos, partidas, registros, gastos


# ── Avance de obra ─────────────────────────────────────────────────────────
def test_avance_por_partidas_es_aporte_al_total_y_sube_a_la_raiz():
    conceptos, partidas, registros, gastos = _datos()
    res = _avance_obra_resumen(conceptos, partidas, registros, gastos, INI, FIN)
    assert res["presupuesto_total"] == Decimal("4000")
    f1, f2 = res["filas"]
    # c1: 200/4000 = 5%; subpartida suma a su raíz: c2 acumulado 400/4000 = 10%, semanal 100/4000 = 2.5%
    assert (f1["partida"], round(f1["acumulado"], 2), round(f1["semanal"], 2)) == ("ESTUDIOS", Decimal("5.00"), Decimal("0.00"))
    assert (f2["partida"], round(f2["acumulado"], 2), round(f2["semanal"], 2)) == ("OBRA CIVIL", Decimal("10.00"), Decimal("2.50"))
    # el avance futuro (después del corte) no cuenta; el total real = suma de partidas
    assert round(res["real_pct"], 2) == round(res["total_acumulado"], 2) == Decimal("15.00")
    assert 0 < res["programado_pct"] < 100


def test_sin_presupuesto_no_hay_resumen():
    assert _avance_obra_resumen([], [], [], [], INI, FIN) is None


def test_registros_sin_partida_no_se_pierden_del_total():
    conceptos, partidas, registros, gastos = _datos()
    registros.append({"concepto_id": "inactivo", "fecha": INI.isoformat(), "importe_ganado_congelado": "400"})
    res = _avance_obra_resumen(conceptos, partidas, registros, gastos, INI, FIN)
    assert res["filas"][-1]["partida"] == "OTRAS PARTIDAS"
    assert round(res["total_acumulado"], 2) == round(res["real_pct"], 2) == Decimal("25.00")


def test_bloques_tabla_con_total_y_resumen():
    conceptos, partidas, registros, gastos = _datos()
    res = _avance_obra_resumen(conceptos, partidas, registros, gastos, INI, FIN)
    resumen, tabla = _avance_obra_bloques(res, INI, FIN)
    assert resumen["type"] == "summary" and "REAL 15.00%" in resumen["text"]
    assert tabla["headers"] == ["N°", "PARTIDAS", "AVANCE SEMANAL (%)", "AVANCE ACUMULADO (%)"]
    assert tabla["rows"][0] == ["1", "ESTUDIOS", "0.00%", "5.00%"]
    assert tabla["rows"][-1] == ["", "TOTAL", "2.50%", "15.00%"]
    assert all(len(r) == 4 for r in tabla["rows"])


def test_graficas_son_png_validas():
    conceptos, partidas, registros, gastos = _datos()
    res = _avance_obra_resumen(conceptos, partidas, registros, gastos, INI, FIN)
    graficas = _avance_obra_graficas(res, FIN)
    assert [g[0] for g in graficas] == ["GRÁFICA DE AVANCE (OBRA)", "FÍSICO – FINANCIERO (OBRA)"]
    from PIL import Image
    for _, _, png in graficas:
        assert Image.open(io.BytesIO(png)).size == (1800, 800)
    assert _avance_obra_graficas({**res, "curva": {"meses": [], "pv": [], "ev": [], "ac": []}}, FIN) == []
    Image.open(io.BytesIO(_grafica_lineas_png("t", ["01/26"], [("a", "#2563EB", [Decimal(5)])], lambda v: f"{v}")))


# ── Estimaciones ───────────────────────────────────────────────────────────
def test_estimacion_valor_normaliza_y_valida():
    assert _estimacion_valor("obra_ejecutada", "$ 5,978,451.30") == "5978451.30"
    assert _estimacion_valor("pct_amortizacion_anticipo", "0.22 %") == "0.22"
    assert _estimacion_valor("estatus", "  LIBERADA ") == "LIBERADA"
    assert _estimacion_valor("iva", "  ") is None
    with pytest.raises(HTTPException):
        _estimacion_valor("periodo_fin", "31/12/2025")
    with pytest.raises(HTTPException):
        _estimacion_valor("iva", "mucho")


def test_estimaciones_bloques_suma_acumulada_calculada_y_un_bloque_por_contrato():
    regs = [
        {"contrato": "obra", "numero": 2, "periodo_inicio": "2026-01-01", "periodo_fin": "2026-01-31", "estatus": "NO LIBERADA",
         "obra_ejecutada": "42492882.52", "iva": "6798861.20", "fecha_pago": None},
        {"contrato": "obra", "numero": 1, "periodo_inicio": "2025-12-08", "periodo_fin": "2025-12-31", "estatus": "NO LIBERADA",
         "obra_ejecutada": "5978451.30", "pct_amortizacion_anticipo": "0.22", "fecha_pago": "Por pagar"},
        {"contrato": "supervision", "numero": 1, "obra_ejecutada": "4480989.20", "estatus": "LIBERADA"},
    ]
    obra, sup = _estimaciones_bloques(regs)
    assert obra["heading"].endswith("(OBRA)") and sup["heading"].endswith("(SUPERVISIÓN)")
    r1, r2 = obra["rows"]
    assert r1[0] == "1" and r1[1] == "08/12/2025 al 31/12/2025" and r1[3] == "$ 5,978,451.30"
    assert r1[4] == "$ 5,978,451.30" and r1[6] == "0.22%" and r1[10] == "Por pagar" or r1[10] == "-"
    assert r2[4] == "$ 48,471,333.82" and r2[10] == "-" and r2[5] == "-"
    assert all(len(r) == len(obra["headers"]) for r in obra["rows"] + sup["rows"])
    assert abs(sum(obra["col_widths"]) - 1.0) < 1e-9
    assert _estimaciones_bloques([]) == []


def test_pendientes_solo_vigentes_por_contrato():
    regs = [
        {"contrato": "obra", "titulo": "Control de Avance", "descripcion": "Pendiente desglose", "orden": 2},
        {"contrato": "obra", "titulo": "Riesgo en pagos", "descripcion": None, "orden": 1},
        {"contrato": "obra", "titulo": "Resuelto", "resuelto": True},
        {"contrato": "supervision", "titulo": "Solo resuelto", "resuelto": True},
    ]
    bloques = _pendientes_bloques(regs)
    assert len(bloques) == 1 and bloques[0]["heading"].endswith("(OBRA)")
    assert [i["title"] for i in bloques[0]["items"]] == ["Riesgo en pagos", "Control de Avance"]
    assert _pendientes_bloques([]) == []


# ── Punta a punta (Mongo simulada) ─────────────────────────────────────────
mongomock_motor = pytest.importorskip("mongomock_motor")

USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Coord", "email": "c@x.mx"}
SUB = {"id": "u2", "role": server.ROLE_SUB, "name": "Sub", "email": "s@x.mx", "project_ids": ["p1"]}
TODOS = {"avance_obra": True, "estimaciones": True, "pendientes_generales": True}


def _db():
    from mongomock_motor import AsyncMongoMockClient
    return AsyncMongoMockClient()["t"]


async def _sembrar(db, modulos):
    await db.projects.insert_one({
        "id": "p1", "name": "Proyecto Demo", "color_institucional": "#003366",
        "modulos": {"reportes_diarios": {}, "reportes_semanales": modulos, "reportes_mensuales": {}},
    })
    conceptos, partidas, registros, gastos = _datos()
    await db.conceptos.insert_many(conceptos)
    await db.partidas.insert_many([{**p, "project_id": "p1"} for p in partidas])
    await db.registros_avance.insert_many([{**r, "project_id": "p1"} for r in registros])
    await db.gastos_reales.insert_many([{**g, "project_id": "p1"} for g in gastos])
    await server.create_estimacion("p1", server.EstimacionIn(
        numero=1, periodo_inicio="2025-12-08", periodo_fin="2025-12-31", estatus="NO LIBERADA",
        obra_ejecutada="5978451.30", total_neto="2483831.29"), SUB)
    await server.create_pendiente_general("p1", server.PendienteIn(
        titulo="Control de Avance", descripcion="Pendiente la entrega del desglose de la Partida 9"), SUB)


def _exportar(formato, apagar=False):
    db = _db()
    old = server.db
    server.db = db
    try:
        async def go():
            await _sembrar(db, TODOS)
            if apagar:
                await db.projects.update_one({"id": "p1"}, {"$set": {"modulos.reportes_semanales": {}}})
            common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=USER)
            if formato == "pdf":
                resp = await server.export_reports_pdf(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None,
                    estructura_solo=False)
            elif formato == "pptx":
                resp = await server.export_reports_pptx(
                    **common, recuento_nivel=None, recuento_node_ids=None, recuento_sin_datos=None)
            else:
                resp = await server.export_reports_docx(**common)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def _texto(formato, data):
    if formato == "pdf":
        from pypdf import PdfReader
        t = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    elif formato == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        partes = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    partes.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    partes += [c.text for r in sh.table.rows for c in r.cells]
        t = "\n".join(partes)
    else:
        from docx import Document
        d = Document(io.BytesIO(data))
        t = "\n".join([p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells])
    return " ".join(t.split())


@pytest.mark.parametrize("formato", ["pdf", "pptx", "docx"])
def test_el_informe_semanal_imprime_la_capa_financiera(formato):
    texto = _texto(formato, _exportar(formato))
    for esperado in (
        "TABLA DE AVANCE POR PARTIDAS (OBRA)", "ESTUDIOS", "OBRA CIVIL", "AVANCE SEMANAL", "TOTAL",
        "15.00%", "REAL",
        "CONCENTRADO ESTIMACIONES (OBRA)", "08/12/2025 al 31/12/2025", "5,978,451.30", "2,483,831.29",
        "PENDIENTES O REQUERIMIENTOS GENERALES (OBRA)", "Control de Avance", "desglose de la Partida 9",
    ):
        assert esperado in texto, f"falta {esperado!r} en el {formato}"


def test_graficas_como_imagenes_en_pptx_y_modulos_apagados_no_imprimen():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx")))
    assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13) >= 2
    texto = _texto("pptx", _exportar("pptx", apagar=True))
    assert "TABLA DE AVANCE POR PARTIDAS" not in texto and "ESTIMACIONES" not in texto and "Control de Avance" not in texto


def test_estimaciones_duplicadas_y_permisos():
    db = _db()
    old = server.db
    server.db = db
    try:
        async def go():
            await _sembrar(db, TODOS)
            with pytest.raises(HTTPException) as e:
                await server.create_estimacion("p1", server.EstimacionIn(numero=1), SUB)
            assert e.value.status_code == 409
            jefe = {"id": "u4", "role": server.ROLE_JEFE, "name": "Jefe", "email": "j@x.mx"}
            with pytest.raises(HTTPException) as e2:
                await server.create_pendiente_general("p1", server.PendienteIn(titulo="x"), jefe)
            assert e2.value.status_code == 403
            lista = await server.list_estimaciones("p1", None, SUB)
            assert len(lista) == 1 and lista[0]["obra_ejecutada"] == "5978451.30"
            upd = await server.update_estimacion("p1", lista[0]["id"], server.EstimacionPatch(estatus="LIBERADA"), SUB)
            assert upd["estatus"] == "LIBERADA"
        asyncio.run(go())
    finally:
        server.db = old
