"""Presentación semanal: plantilla "semanal", selector de reportes de la semana
y export PPTX con las láminas del informe de referencia.

Usa el Mongo local (DB de prueba) y llama directamente a las funciones de los
endpoints — sin HTTP ni auth real."""
import asyncio
import base64
import io
import os
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_presentacion_semanal")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation  # noqa: E402

import server  # noqa: E402

# PNG 1x1 válido
PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="


def _run(coro):
    return asyncio.run(coro)


def _template_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = 9144000 * 3, 6858000 * 3
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))


async def _scenario(tmp_path: Path, con_plantilla: bool):
    pid, uid = str(uuid.uuid4()), str(uuid.uuid4())
    user = {"id": uid, "name": "Coord", "role": server.ROLE_COORD, "project_ids": [pid]}
    proj = {"id": pid, "name": "Proyecto X", "objeto_contrato": "OBJETO DEL CONTRATO",
            "color_institucional": "#8A1538",
            "modulos": {"reportes_semanales": {"impacto_ambiental": True, "aspectos_sociales": True,
                                                "seguridad_higiene": True, "estudios_proyecto": True,
                                                "pendientes_generales": True}},
            "datos_contrato": {"obra": {"empresa": "ALFA", "contrato": "C-1", "monto": "$1.00"}}}
    if con_plantilla:
        tpl = tmp_path / "tpl.pptx"
        _template_pptx(tpl)
        proj["template_pptx"] = str(tpl)
    await server.db.projects.insert_one(proj)
    now = datetime.now(timezone.utc)
    rid_in, rid_out = str(uuid.uuid4()), str(uuid.uuid4())
    base = {"project_id": pid, "node_id": "n", "node_path_names": ["Tramo", "Estación 1"],
            "images": [PNG_B64] * 5, "photo_captions": ["uno", "dos"], "actividades": "Colado\nHabilitado",
            "severidad": ["informativo"], "captured_by_name": "Esp"}
    await server.db.reports.insert_many([
        {**base, "id": rid_in, "title": "Colado de zapatas", "created_at": now},
        {**base, "id": rid_out, "title": "Reporte viejo", "created_at": datetime(2020, 1, 1, tzinfo=timezone.utc)},
    ])
    await server.db.ambiental_auditoria.insert_one({
        "id": str(uuid.uuid4()), "project_id": pid, "tipo": "cumplimiento", "descripcion": "Riego de areas verdes",
        "fecha": now.date().isoformat(), "created_at": now, "captured_by_name": "Esp", "images": []})
    await server.db.levantamientos_sociales.insert_one({
        "id": str(uuid.uuid4()), "project_id": pid, "tipo": "dictamen_estructural", "ambito": "estacion",
        "total_requerido": 10, "realizados": 5})
    await server.db.seguridad_reportes.insert_one({
        "id": str(uuid.uuid4()), "project_id": pid, "fecha": now.date().isoformat(), "created_at": now,
        "trabajador": "Juan Perez", "contratista": "ALFA", "descripcion": "Caida leve", "severidad": "leve", "images": []})
    await server.db.estudios_bitacora.insert_one({
        "id": str(uuid.uuid4()), "project_id": pid, "disciplina": "topografia",
        "texto": "Levantamiento de poligonal", "created_at": now})
    await server.db.pendientes_generales.insert_one({
        "id": str(uuid.uuid4()), "project_id": pid, "contrato": "obra", "titulo": "Control de Avance",
        "descripcion": "Programa pendiente de aval", "resuelto": False, "orden": 0, "created_at": now})
    try:
        fecha = now.date().isoformat()
        created = await server.create_presentacion(
            pid, server.PresentacionIn(project_id=pid, fecha=fecha, plantilla="semanal"), user)
        tipos = [s["tipo"] for s in sorted(created["secciones"], key=lambda s: s["orden"])]
        assert tipos[0] == "portada" and tipos[-1] == "cierre" and "reportes_semana" in tipos

        semana = await server.list_reportes_semana_presentacion(pid, created["id"], user)
        ids = [r["id"] for r in semana]
        assert rid_in in ids and rid_out not in ids
        assert semana[0]["titulo"] == "Colado de zapatas" and semana[0]["n_fotos"] == 5

        sec = next(s for s in created["secciones"] if s["tipo"] == "reportes_semana")
        await server.update_presentacion_seccion(
            pid, created["id"], sec["id"], server.PresentacionSeccionPatch(report_ids=[rid_in]), user)

        resp = await server.export_presentacion_pptx(pid, created["id"], user)
        data = b"".join([chunk async for chunk in resp.body_iterator])
        return Presentation(io.BytesIO(data)), tipos
    finally:
        await server.db.projects.delete_one({"id": pid})
        await server.db.presentaciones.delete_many({"project_id": pid})
        await server.db.reports.delete_many({"project_id": pid})
        for col in ("ambiental_auditoria", "levantamientos_sociales", "seguridad_reportes",
                    "estudios_bitacora", "pendientes_generales"):
            await server.db[col].delete_many({"project_id": pid})


async def _ambos(tmp_path):
    # Un solo bucle de eventos: el cliente Motor queda atado al primero.
    return [await _scenario(tmp_path, False), await _scenario(tmp_path, True)]


def test_export_semanal_incluye_reporte_elegido(tmp_path):
    for prs, _tipos in _run(_ambos(tmp_path)):
        _verifica(prs)


def _verifica(prs):
    textos = []
    for sl in prs.slides:
        partes = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        for sh in sl.shapes:
            if sh.has_table:
                partes += [c.text for row in sh.table.rows for c in row.cells]
        textos.append(" ".join(partes))
    todo = " | ".join(textos)
    assert "INFORME DE ACTIVIDADES" in todo
    assert "Colado de zapatas" in todo           # reporte elegido
    assert "Reporte viejo" not in todo           # fuera de la semana / no elegido
    assert "(cont.)" in todo                     # 5 fotos → segunda lámina
    assert "Gracias" in todo
    assert "Riego de areas verdes" in todo       # Ambiental automático
    assert "AMBIENTAL" in todo and "SEGURIDAD" in todo and "SOCIAL / IMPACTO URBANO" in todo
    assert "AVANCE EN DICTÁMENES" in todo        # Social automático
    assert "Levantamiento de poligonal" in todo  # Estudios y proyecto ejecutivo
    assert "Programa pendiente de aval" in todo  # Pendientes generales
    assert "Juan Perez" in todo                  # Seguridad automática
    assert "DATOS DE CONTRATO (OBRA)" in todo    # financiero con datos de contrato
    # portada, mapa, 3 divisores, contrato, estudios, 2 láminas de reporte, cierre
    assert len(prs.slides) >= 10
