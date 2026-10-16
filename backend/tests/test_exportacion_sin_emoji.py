"""Los emoji (semáforo 🟢🟡🔴, etc.) son código interno: no salen en las
exportaciones, pero se conservan en los datos y en la app."""
import os
import sys
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_emoji")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation  # noqa: E402
from pptx.util import Cm  # noqa: E402

from server import (_quitar_emoji, _quitar_emoji_deep, _quitar_emoji_pptx,  # noqa: E402
                    _quitar_emoji_reportes)


def test_quita_emoji_y_el_espacio_pegado():
    assert _quitar_emoji("🟢 Se ejecuta excavación en pozo PVP-02.") == "Se ejecuta excavación en pozo PVP-02."
    assert _quitar_emoji("🟠Se ejecuta el acomodo") == "Se ejecuta el acomodo"
    assert _quitar_emoji("Actividades:🔴 revisión") == "Actividades: revisión"


def test_no_toca_simbolos_de_texto_ni_acentos():
    t = "Estación 3 › Ñandú — 5° · ✓ → ☐ • 2²"
    assert _quitar_emoji(t) == t
    assert _quitar_emoji(None) is None and _quitar_emoji(7) == 7


def test_deep_limpia_estructuras_anidadas():
    datos = {"a": ["🟢 uno", {"b": "🟡 dos"}], "n": 3, "t": ("🔴 tres",)}
    assert _quitar_emoji_deep(datos) == {"a": ["uno", {"b": "dos"}], "n": 3, "t": ("tres",)}


def test_reportes_solo_texto_y_no_imagenes():
    rep = [{"actividades": "🟢 Colado", "images": ["data:image/png;base64,AAA🟢"], "dynamic_data": {"x": "🔴 y"}}]
    _quitar_emoji_reportes(rep)
    assert rep[0]["actividades"] == "Colado" and rep[0]["dynamic_data"] == {"x": "y"}
    assert rep[0]["images"] == ["data:image/png;base64,AAA🟢"]  # no se toca


def test_pptx_ya_armado_sin_emoji():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(0, 0, Cm(5), Cm(1))
    tb.text_frame.text = "🟢 Se ejecuta excavación"
    _quitar_emoji_pptx(prs)
    assert tb.text_frame.text == "Se ejecuta excavación"
