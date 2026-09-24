"""A-09: el filtro por partida en Visualizar avance incluye a TODA la descendencia (3 niveles)
en el listado, el resumen y los exports PDF y XLSX. Pruebas sin Mongo (db falso)."""
import asyncio
import io
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_avance_filtro_partida")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402

import server  # noqa: E402

PID = "p1"
USER = {"id": "u1", "role": server.ROLE_COORD, "name": "Ana"}

# Árbol de partidas (3 niveles) y un concepto con avance en cada una:
#   raiz (c-raiz)
#     └─ medio (c-medio)
#          └─ hoja (c-hoja)
#   otra-raiz (c-otra)
PARTIDAS = [
    {"id": "raiz", "project_id": PID, "partida_padre_id": None},
    {"id": "medio", "project_id": PID, "partida_padre_id": "raiz"},
    {"id": "hoja", "project_id": PID, "partida_padre_id": "medio"},
    {"id": "otra-raiz", "project_id": PID, "partida_padre_id": None},
]
CONCEPTOS = [
    {"id": f"c-{p}", "project_id": PID, "partida_id": p, "clave": f"K-{p}", "descripcion": f"Concepto {p}",
     "unidad": "M2", "cantidad": "100", "precio_unitario": "10", "meta": "", "periodicidad_meta": "SEMANAL", "activo": True}
    for p in ("raiz", "medio", "hoja", "otra-raiz")
]
AVANCES = [
    {"id": f"a-{p}", "project_id": PID, "concepto_id": f"c-{p}", "fecha": "2026-09-20", "cantidad_ejecutada": "5",
     "capturado_por": "u1", "creado_en": datetime(2026, 9, 20, tzinfo=timezone.utc)}
    for p in ("raiz", "medio", "hoja", "otra-raiz")
]


def _coincide(doc, filtro):
    for campo, cond in filtro.items():
        valor = doc.get(campo)
        if isinstance(cond, dict):
            if "$in" in cond and valor not in cond["$in"]:
                return False
            if "$gte" in cond and not (valor is not None and valor >= cond["$gte"]):
                return False
            if "$lte" in cond and not (valor is not None and valor <= cond["$lte"]):
                return False
        elif valor != cond:
            return False
    return True


class _Cursor:
    def __init__(self, docs):
        self.docs = docs

    def sort(self, *_a, **_k):
        return self

    async def to_list(self, length=None):
        return [dict(d) for d in self.docs]


class _Coleccion:
    def __init__(self, docs):
        self.docs = docs
        self.consultas = []

    def find(self, filtro=None, *_a, **_k):
        self.consultas.append(filtro or {})
        return _Cursor([d for d in self.docs if _coincide(d, filtro or {})])

    async def find_one(self, filtro, *_a, **_k):
        for d in self.docs:
            if _coincide(d, filtro):
                return dict(d)
        return None


class _Db:
    def __init__(self):
        self.partidas = _Coleccion(PARTIDAS)
        self.conceptos = _Coleccion(CONCEPTOS)
        self.registros_avance = _Coleccion(AVANCES)
        self.location_nodes = _Coleccion([])
        self.projects = _Coleccion([{"id": PID, "name": "Proyecto de prueba"}])


@pytest.fixture()
def db(monkeypatch):
    falso = _Db()
    monkeypatch.setattr(server, "db", falso)

    async def acceso(user, pid):
        return {"id": pid, "name": "Proyecto de prueba"}

    monkeypatch.setattr(server, "ensure_project_access", acceso)
    return falso


def _correr(coro):
    return asyncio.run(coro)


def _ids_de_avances(salida):
    return {a.concepto_id for a in salida}


def test_raiz_incluye_hijas_y_nietas_pero_no_otra_rama(db):
    salida = _correr(server.list_avances(PID, partida_id="raiz", user=USER))
    assert _ids_de_avances(salida) == {"c-raiz", "c-medio", "c-hoja"}


def test_un_nivel_intermedio_incluye_sus_descendientes_pero_no_a_su_padre(db):
    salida = _correr(server.list_avances(PID, partida_id="medio", user=USER))
    assert _ids_de_avances(salida) == {"c-medio", "c-hoja"}


def test_una_hoja_solo_se_incluye_a_si_misma(db):
    salida = _correr(server.list_avances(PID, partida_id="hoja", user=USER))
    assert _ids_de_avances(salida) == {"c-hoja"}


def test_sin_filtro_trae_todo(db):
    assert len(_correr(server.list_avances(PID, user=USER))) == 4


def test_resumen_suma_toda_la_descendencia(db):
    # 5 x P.U. 10 = 50 por registro; raíz + medio + hoja = 150.
    resumen = _correr(server.avances_resumen(PID, partida_id="raiz", user=USER))
    assert resumen.total_registros == 3
    assert resumen.importe_total == "150.00"
    medio = _correr(server.avances_resumen(PID, partida_id="medio", user=USER))
    assert (medio.total_registros, medio.importe_total) == (2, "100.00")


def _bytes_de(respuesta):
    if hasattr(respuesta, "body") and respuesta.body:
        return respuesta.body
    async def _leer():
        partes = []
        async for parte in respuesta.body_iterator:
            partes.append(parte if isinstance(parte, bytes) else parte.encode())
        return b"".join(partes)
    return _correr(_leer())


def test_export_xlsx_incluye_toda_la_descendencia(db):
    from openpyxl import load_workbook

    respuesta = _correr(server.export_avance_periodo_xlsx(PID, partida_id="raiz", user=USER))
    hoja = load_workbook(io.BytesIO(_bytes_de(respuesta))).active
    textos = [str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None]
    claves = {t.split(" — ")[0] for t in textos if t.startswith("K-")}
    assert claves == {"K-raiz", "K-medio", "K-hoja"}


def test_export_pdf_filtra_con_el_mismo_subtree(db):
    respuesta = _correr(server.export_avance_periodo_pdf(PID, partida_id="medio", user=USER))
    assert getattr(respuesta, "status_code", 200) == 200
    consultas = [q for q in db.registros_avance.consultas if "concepto_id" in q]
    assert consultas and set(consultas[-1]["concepto_id"]["$in"]) == {"c-medio", "c-hoja"}
