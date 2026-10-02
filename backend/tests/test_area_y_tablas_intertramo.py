"""Portadilla de Área con el nombre REAL del área (antes un Sub-coordinador veía
"OBRA CIVIL" en todas), título de Estación sin duplicar y tabla de postes del
resumen de intertramo con letra grande y filas que crecen con su texto (antes
un "Cambio de trayectoria" se encimaba con la fila siguiente)."""
import asyncio
import io
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_area_tablas")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402

import server  # noqa: E402
from server import ROLE_COORD, ROLE_SUB, _estacion_titulos, _normalize_area_name  # noqa: E402


# ── Regla de área ──────────────────────────────────────────────────────────
def test_un_subcoordinador_ve_el_area_real_de_cada_reporte():
    assert _normalize_area_name("Seguridad", user_role=ROLE_SUB, user_area=None) == "Seguridad"
    assert _normalize_area_name("CALIDAD", user_role=ROLE_SUB, user_area="Topografía") == "Calidad"  # no la pisa con la suya
    assert _normalize_area_name("Seguridad", user_role=ROLE_COORD) == "Seguridad"


def test_nombre_generico_historico_y_area_vacia_siguen_igual():
    for generico in ("Sub supervisores", "sub-supervisor", "SUBSUPERVISORES"):
        assert _normalize_area_name(generico, user_role=ROLE_SUB, user_area="Topografía") == "TOPOGRAFÍA"
        assert _normalize_area_name(generico, user_role=ROLE_SUB, user_area=None) == "OBRA CIVIL"
        assert _normalize_area_name(generico, user_role=ROLE_COORD) == "OBRA CIVIL"
    assert _normalize_area_name("", user_role=ROLE_SUB, user_area="Calidad") == "CALIDAD"  # sin área: la del sub
    assert _normalize_area_name("", user_role=ROLE_COORD) == "SIN ÁREA"
    assert _normalize_area_name("calidad", user_role=ROLE_COORD) == _normalize_area_name("Calidad", user_role=ROLE_COORD)


# ── Título de Estación ─────────────────────────────────────────────────────
def test_titulo_de_estacion_no_se_imprime_dos_veces():
    assert _estacion_titulos(1, "ESTACIÓN 1. UNIVERSIDAD") == ("ESTACIÓN 1.", "UNIVERSIDAD")
    assert _estacion_titulos(5, "Estación 5 CEFORMA") == ("ESTACIÓN 5.", "CEFORMA")
    assert _estacion_titulos(7, "Estación 7") == (None, "Estación 7")  # el nombre ES el ordinal: solo el grande
    assert _estacion_titulos(2, "CANTERA") == ("ESTACIÓN 2.", "CANTERA")
    assert _estacion_titulos(3, "Estación 8 OTRA") == ("ESTACIÓN 3.", "Estación 8 OTRA")  # otro número: se respeta
    assert _estacion_titulos(None, "Estación 1") == (None, "Estación 1")


# ── Exportes (Mongo simulada) ──────────────────────────────────────────────
pytest.importorskip("mongomock_motor")
SUB = {"id": "u2", "role": ROLE_SUB, "name": "Sub", "email": "s@x.mx", "project_ids": ["p1"]}
LARGO = "Cambio de trayectoria"


async def _sembrar(db, n_postes):
    await db.projects.insert_one({"id": "p1", "name": "Proyecto Demo", "color_institucional": "#7A1030",
                                  "modulos": {"reportes_diarios": {}, "reportes_semanales": {}, "reportes_mensuales": {}}})
    nodos = [("r1", "TRAMO 1", None, 0, False, None), ("e1", "Estación 1", "r1", 1, True, None),
             ("i12", "Intertramo 1-2", "r1", 1, False, "intertramo")]
    nodos += [(f"s{i}", f"S{i}", "i12", 2, True, None) for i in range(n_postes)]
    for i, (nid, name, parent, depth, leaf, tipo) in enumerate(nodos):
        doc = {"id": nid, "project_id": "p1", "name": name, "parent_id": parent, "depth": depth, "order": i,
               "is_leaf": leaf, "measurement_type": "cadenamiento"}
        if tipo:
            doc["tipo_nodo"] = tipo
        if nid.startswith("s"):
            doc.update({"ubicacion": "Av. Imán", "estatus_pilas": "cambio_trayectoria" if i % 2 else "sin_iniciar",
                        "estatus_cabezal": "cambio_trayectoria", "estatus_fuste": "concluido", "estatus_poste": "en_proceso"})
        await db.location_nodes.insert_one(doc)
    await db.areas.insert_one({"id": "a1", "project_id": "p1", "name": "SEGURIDAD"})
    await db.reports.insert_one({
        "id": "r1", "project_id": "p1", "node_id": "e1", "area_id": "a1", "created_at": datetime.now(timezone.utc),
        "severidad": "diario", "actividades": "Recorrido de seguridad e higiene", "images": [], "personnel": [],
        "equipment": [], "captured_by": "u2", "captured_by_name": "Ing", "node_path_names": ["TRAMO 1", "Estación 1"],
        "title": "Recorrido", "measurement_value": {"cadenamiento": "1+000"}})


def _exportar(formato, user, n_postes=20):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await _sembrar(db, n_postes)
            common = dict(pid="p1", period="today", period_type="diario", area_id=None, scope=None, user=user)
            if formato == "pdf":
                resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                       recuento_sin_datos=None, estructura_solo=False)
            else:
                resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None,
                                                        recuento_sin_datos=None)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def test_portadilla_pdf_dice_el_area_real_para_un_subcoordinador():
    from pypdf import PdfReader
    paginas = [p.extract_text() or "" for p in PdfReader(io.BytesIO(_exportar("pdf", SUB, 3))).pages]
    texto = "\n".join(paginas)
    assert any(t.strip().upper().startswith("SEGURIDAD") and len(t.strip()) < 40 for t in paginas), "portadilla del área real"
    assert "OBRA CIVIL" not in texto


def test_portadilla_pptx_dice_el_area_real_para_un_subcoordinador():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx", SUB, 3)))
    textos = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]
    assert any(t.strip().upper() == "SEGURIDAD" for t in textos)
    assert not any("OBRA CIVIL" in t for t in textos)


def _palabras_encimadas(pagina) -> list:
    """Pares de palabras cuyas cajas se traslapan de verdad (texto encimado)."""
    ws = pagina.get_text("words")
    malos = []
    for i in range(len(ws)):
        x0, y0, x1, y1 = ws[i][:4]
        for j in range(i + 1, len(ws)):
            a0, b0, a1, b1 = ws[j][:4]
            ancho, alto = min(x1, a1) - max(x0, a0), min(y1, b1) - max(y0, b0)
            if ancho > 1.5 and alto > 2.5:  # traslape real, no roce de bordes
                malos.append((ws[i][4], ws[j][4]))
    return malos


@pytest.mark.parametrize("n_postes", [4, 20, 45])
def test_pdf_intertramo_sin_texto_encimado_y_con_letra_grande(n_postes):
    fitz = pytest.importorskip("pymupdf")
    doc = fitz.open(stream=_exportar("pdf", SUB, n_postes), filetype="pdf")
    paginas = [p for p in doc if "ACTIVIDAD EN POSTES" in p.get_text().upper()]
    assert paginas, "debe existir la página de resumen del intertramo"
    for pag in paginas:
        assert not _palabras_encimadas(pag), f"texto encimado en la página {pag.number + 1}"
        assert LARGO.split()[0] in pag.get_text()  # "Cambio" sale completo
    # con pocos postes la letra de la tabla es grande (≥ 9 pt); con muchos baja pero no se encima
    tam = [s["size"] for b in paginas[0].get_text("dict")["blocks"] for l in b.get("lines", []) for s in l["spans"]
           if s["text"].strip() in ("Sin iniciar", "Concluido", "En proceso")]
    assert tam, "se esperaban celdas de estatus"
    if n_postes <= 4:
        assert max(tam) >= 9.5
    if n_postes == 45:
        assert "más…" in paginas[0].get_text()


def test_pptx_intertramo_filas_crecen_y_nada_se_sale_de_la_diapositiva():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx", SUB, 10)))
    fuera = [(i + 1, sh.name) for i, sl in enumerate(prs.slides) for sh in sl.shapes
             if sh.top is not None and sh.top + sh.height > prs.slide_height + 5_000]
    assert not fuera
    textos = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]
    pagina = next(t for t in textos if "ACTIVIDAD EN POSTES" in t.upper())
    # con 10 postes los estatus largos se abrevian a una línea; ninguno se recorta con "+N más"
    assert "Cambio tray." in pagina and "más…" not in pagina


def test_pptx_intertramo_con_pocos_postes_usa_el_texto_completo():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(_exportar("pptx", SUB, 4)))
    textos = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]
    pagina = next(t for t in textos if "ACTIVIDAD EN POSTES" in t.upper())
    assert LARGO in pagina and "Cambio tray." not in pagina
