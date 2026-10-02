"""Relación de juntas de trabajo (calculada de las minutas) y hoja final del
informe con el superintendente de supervisión (Datos de contrato > Supervisión)."""
import asyncio
import io
import os
import sys
from datetime import date, datetime, timezone
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_juntas_cierre")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402

import server  # noqa: E402
from server import ROLE_COORD, _cierre_firma, _junta_consecutivo, _juntas_bloques, _periodo_texto_cierre  # noqa: E402

CARGO = "SUPERINTENDENTE DE SUPERVISIÓN EXTERNA"


# ── Puras ──────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("codigo,esperado", [
    ("CL4 OBRA 260921-A", "CL4-OBRA-260921A"), ("CL4 OBRA 260923-AB", "CL4-OBRA-260923AB"),
    ("260921-C", "260921C"), ("texto libre", "texto libre"), ("", ""), (None, ""),
])
def test_consecutivo_de_junta(codigo, esperado):
    assert _junta_consecutivo(codigo) == esperado


def test_bloque_de_juntas_ordenado_y_formateado():
    minutas = [
        {"codigo": "CL4 OBRA 260923-A", "fecha_reunion": "2026-09-23", "titulo": "Análisis de movilidad", "created_at": 2},
        {"codigo": "CL4 OBRA 260921-B", "fecha_reunion": "2026-09-21", "titulo": "Mesa de la Estimación 9", "created_at": 1},
        {"codigo": "CL4 OBRA 260921-A", "fecha_reunion": "2026-09-21", "titulo": "Salvamento arqueológico", "created_at": 0},
    ]
    (b,) = _juntas_bloques(minutas)
    assert b["headers"] == ["N°", "CONSECUTIVO", "FECHA", "ASUNTO"]
    assert b["rows"] == [
        ["1", "CL4-OBRA-260921A", "21-sep-26", "Salvamento arqueológico"],
        ["2", "CL4-OBRA-260921B", "21-sep-26", "Mesa de la Estimación 9"],
        ["3", "CL4-OBRA-260923A", "23-sep-26", "Análisis de movilidad"],
    ]
    assert _juntas_bloques([]) == []


@pytest.mark.parametrize("ini,fin,texto", [
    (date(2026, 9, 21), date(2026, 9, 26), "DEL 21 AL 26 DE SEPTIEMBRE DEL 2026"),
    (date(2026, 9, 28), date(2026, 10, 3), "DEL 28 DE SEPTIEMBRE AL 03 DE OCTUBRE DEL 2026"),
    (date(2026, 12, 28), date(2027, 1, 2), "DEL 28 DE DICIEMBRE DE 2026 AL 02 DE ENERO DE 2027"),
    (date(2026, 9, 21), date(2026, 9, 21), "21 DE SEPTIEMBRE DEL 2026"),
])
def test_periodo_texto_cierre(ini, fin, texto):
    assert _periodo_texto_cierre(ini, fin) == texto


def _dt(y, m, d, h=12):
    return datetime(y, m, d, h, 0, tzinfo=server.MX_TZ).astimezone(timezone.utc)


def test_firma_solo_si_hay_nombre():
    ini, fin = _dt(2026, 9, 21), _dt(2026, 9, 26)
    assert _cierre_firma({}, "reportes_semanales", ini, fin) is None
    assert _cierre_firma({"datos_contrato": {"supervision": {"objeto": "X"}}}, "reportes_semanales", ini, fin) is None
    f = _cierre_firma({"datos_contrato": {"supervision": {"superintendente": " Ing. Jorge Ramírez Mejía "}}},
                      "reportes_semanales", ini, fin)
    assert f == {"titulo": "REPORTE SEMANAL", "periodo": "PERIODO DEL 21 AL 26 DE SEPTIEMBRE DEL 2026",
                 "nombre": "Ing. Jorge Ramírez Mejía", "cargo": CARGO}


def test_superintendente_no_crea_bloque_de_supervision_vacio():
    solo_nombre = {"datos_contrato": {"supervision": {"superintendente": "Ing. X"}}, "name": "P"}
    assert not any(b["etiqueta"] == "SUPERVISIÓN" for b in server._contrato_bloques(solo_nombre))
    # y se conserva al sanear
    assert server._sanitize_datos_contrato(solo_nombre["datos_contrato"])["supervision"]["superintendente"] == "Ing. X"


# ── Exportes (Mongo simulada) ──────────────────────────────────────────────
pytest.importorskip("mongomock_motor")
COORD = {"id": "u1", "role": ROLE_COORD, "name": "C", "email": "c@x.mx", "project_ids": ["p1"]}


async def _sembrar(db, con_nombre=True, con_modulo=True):
    hoy = datetime.now(server.MX_TZ).date().isoformat()
    proj = {"id": "p1", "name": "Proyecto Demo", "color_institucional": "#7A1030",
            "modulos": {"reportes_diarios": {}, "reportes_mensuales": {},
                        "reportes_semanales": {"juntas_trabajo": con_modulo}}}
    if con_nombre:
        proj["datos_contrato"] = {"supervision": {"superintendente": "Ing. Jorge Ramírez Mejía"}}
    await db.projects.insert_one(proj)
    await db.minutas.insert_many([
        {"id": "m1", "project_id": "p1", "codigo": "CL4 OBRA 260921-A", "titulo": "Seguimiento a temas de salvamento arqueológico",
         "fecha_reunion": hoy, "created_at": datetime.now(timezone.utc)},
        {"id": "m2", "project_id": "p1", "codigo": "CL4 OBRA 260921-B", "titulo": "Mesa de trabajo Estimación 9",
         "fecha_reunion": hoy, "created_at": datetime.now(timezone.utc)},
        {"id": "m3", "project_id": "p1", "codigo": "CL4 OBRA 250101-A", "titulo": "Junta de otro periodo",
         "fecha_reunion": "2025-01-01", "created_at": datetime.now(timezone.utc)},
    ])


def _exportar(formato, **kw):
    from mongomock_motor import AsyncMongoMockClient
    db = AsyncMongoMockClient()["t"]
    old = server.db
    server.db = db
    try:
        async def go():
            await _sembrar(db, **kw)
            common = dict(pid="p1", period="week", period_type="semanal", area_id=None, scope=None, user=COORD)
            if formato == "pdf":
                resp = await server.export_reports_pdf(**common, recuento_nivel=None, recuento_node_ids=None,
                                                       recuento_sin_datos=None, estructura_solo=False)
            elif formato == "docx":
                resp = await server.export_reports_docx(**common)
            else:
                resp = await server.export_reports_pptx(**common, recuento_nivel=None, recuento_node_ids=None,
                                                        recuento_sin_datos=None)
            return b"".join([c async for c in resp.body_iterator])
        return asyncio.run(go())
    finally:
        server.db = old


def _textos_pptx(datos):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(datos))
    out = []
    for sl in prs.slides:
        partes = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        partes += [c.text for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table
                   for r in sh.table.rows for c in r.cells]
        out.append("\n".join(partes))
    return out


def test_pptx_juntas_y_hoja_final():
    t = _textos_pptx(_exportar("pptx"))
    juntas = next(x for x in t if "RELACIÓN DE JUNTAS DE TRABAJO" in x and "CONSECUTIVO" in x)
    assert "CL4-OBRA-260921A" in juntas and "CL4-OBRA-260921B" in juntas
    assert "Seguimiento a temas de salvamento arqueológico" in juntas
    assert "Junta de otro periodo" not in juntas and "250101" not in juntas
    ultima = t[-1]
    assert "REPORTE SEMANAL" in ultima and "Ing. Jorge Ramírez Mejía" in ultima and CARGO in ultima
    assert "PERIODO DEL" in ultima


def test_pptx_sin_nombre_ni_modulo_no_imprime_nada_de_eso():
    t = _textos_pptx(_exportar("pptx", con_nombre=False, con_modulo=False))
    assert not any("RELACIÓN DE JUNTAS" in x or CARGO in x for x in t)


def test_pdf_juntas_y_hoja_final():
    from pypdf import PdfReader
    paginas = [p.extract_text() or "" for p in PdfReader(io.BytesIO(_exportar("pdf"))).pages]
    assert any("RELACIÓN DE JUNTAS DE TRABAJO" in x and "CL4-OBRA-260921A" in x for x in paginas)
    assert "Ing. Jorge Ramírez Mejía" in paginas[-1] and CARGO in paginas[-1]


def test_docx_juntas_y_hoja_final():
    from docx import Document
    d = Document(io.BytesIO(_exportar("docx")))
    parrafos = [p.text for p in d.paragraphs]
    celdas = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "RELACIÓN DE JUNTAS DE TRABAJO" in parrafos and "CL4-OBRA-260921B" in celdas
    assert parrafos[-1] == CARGO and "Ing. Jorge Ramírez Mejía" in parrafos
