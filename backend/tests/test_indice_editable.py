"""Tests del índice editable de los exports PDF/PPTX (server.py).

Puros: sin HTTP ni Mongo. Se fabrican un PDF y un PPTX sintéticos con
secciones conocidas y se verifica sobre el documento RESULTANTE que (a) las
secciones quedan en el orden configurado, (b) el número de página que imprime
el índice es exactamente donde cayó cada sección y (c) diario/semanal/mensual
no se contaminan entre sí.
"""
import io
import os
import sys
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "synco_test_indice")
os.environ.setdefault("JWT_SECRET", "dummy_test_secret")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from reportlab.lib.pagesizes import A4, landscape  # noqa: E402
from reportlab.pdfgen import canvas as _canvas  # noqa: E402

import server  # noqa: E402
from server import (  # noqa: E402
    _indice_apply_order, _indice_page_counts, _indice_plan, _indice_ranges,
    _indice_truncate, _pdf_apply_indice, _pptx_apply_indice,
    _sanitize_indice_config,
)


def cfg(**kw):
    return _sanitize_indice_config(kw)


# ── Saneado y orden ────────────────────────────────────────────────────────
def test_sanitize_tolera_basura_y_da_defaults():
    c = _sanitize_indice_config("no soy un dict")
    assert c == {"enabled": True, "titulo": "ÍNDICE", "order": [], "titles": {}, "hidden": []}
    c = _sanitize_indice_config({
        "enabled": "si", "titulo": "  Contenido   del   informe ", "order": ["a", "a", 3, "b"],
        "titles": {"a": "  Uno  ", "b": "   ", 7: "x"}, "hidden": "no-lista",
    })
    assert c["enabled"] is True  # no-bool → default
    assert c["titulo"] == "Contenido del informe"
    assert c["order"] == ["a", "b"]
    assert c["titles"] == {"a": "Uno"}
    assert c["hidden"] == []


def test_apply_order_reparte_en_las_posiciones_ocupadas_y_respeta_nuevas():
    natural = ["gd", "n1", "n2", "n3", "rec"]
    assert _indice_apply_order(natural, []) == natural
    assert _indice_apply_order(natural, ["rec", "n3", "n2", "n1", "gd"]) == ["rec", "n3", "n2", "n1", "gd"]
    # Orden parcial + clave desaparecida ("viejo") + sección nueva ("n3"):
    # las guardadas se permutan entre sus huecos, la nueva no se mueve.
    assert _indice_apply_order(natural, ["n2", "viejo", "gd"]) == ["n2", "n1", "gd", "n3", "rec"]


def test_ranges_descarta_vacias_y_desambigua_duplicadas():
    marks = [("a", "A", 1), ("vacia", "V", 3), ("b", "B", 3), ("a", "A2", 5)]
    r = _indice_ranges(marks, 8)
    assert [(x["key"], x["start"], x["n"]) for x in r] == [("a", 1, 2), ("b", 3, 2), ("a#2", 5, 3)]


def test_page_counts_primera_pagina_mas_corta():
    assert _indice_page_counts(0, 5, 8) == []
    assert _indice_page_counts(5, 5, 8) == [5]
    assert _indice_page_counts(6, 5, 8) == [5, 1]
    assert _indice_page_counts(20, 5, 8) == [5, 8, 7]


def test_truncate_agrega_elipsis_y_respeta_el_limite():
    t = _indice_truncate("x" * 100, lambda s: len(s) <= 20)
    assert len(t) <= 20 and t.endswith("…")
    assert _indice_truncate("corto", lambda s: len(s) <= 20) == "corto"


def test_plan_numera_paginas_despues_de_portada_e_indice():
    ranges = [
        {"key": "a", "title": "A", "start": 1, "n": 2},
        {"key": "b", "title": "B", "start": 3, "n": 1},
        {"key": "c", "title": "C", "start": 4, "n": 3},
    ]
    plan = _indice_plan(ranges, cfg(order=["c", "b", "a"], hidden=["b"], titles={"c": "Cierre"}), 1, 10, 10)
    # portada (1) + 1 pág. de índice (2) → la 1ª sección empieza en la 3.
    assert [r["key"] for r in plan["ordered"]] == ["c", "a"]  # b oculta: no se imprime
    assert plan["entries"] == [("Cierre", 3), ("A", 6)]
    assert plan["counts"] == [2]
    assert _indice_plan(ranges, cfg(enabled=False), 1, 10, 10)["entries"] == []


# ── PDF ────────────────────────────────────────────────────────────────────
PAGE = landscape(A4)
MARKS_PDF = [("gd", "Datos generales", 1), ("node:1", "Tramo 1 › E1", 2), ("node:2", "Tramo 1 › E2", 4),
             ("recuento", "Recuento general", 5)]


def _pdf_sintetico(n_pages=7):
    buf = io.BytesIO()
    c = _canvas.Canvas(buf, pagesize=PAGE)
    for i in range(n_pages):
        c.drawString(100, 300, f"CONTENIDO-{i}")
        c.showPage()
    c.save()
    return buf.getvalue()


def _textos(pdf_bytes):
    return [p.extract_text() for p in PdfReader(io.BytesIO(pdf_bytes)).pages]


def _apply(pdf, marks, conf, tpl=None):
    return _pdf_apply_indice(
        pdf, marks, conf, template_path=tpl, brand_hex="#003366", subtitulo="Reporte diario · 01-01-2026",
        top_margin=113.4, bottom_margin=56.7, margin_x=56.7,
    )


def _entradas(indice_text):
    """{título: nº} leído del texto de la página de índice."""
    # pypdf extrae título y número como líneas consecutivas (x distintas).
    out = {}
    lines = [ln.strip() for ln in indice_text.splitlines() if ln.strip()]
    for prev, ln in zip(lines, lines[1:]):
        if ln.isdigit():
            out[prev] = int(ln)
    return out


def test_pdf_indice_default_apunta_a_las_paginas_reales():
    pdf = _pdf_sintetico(7)
    out, estructura = _apply(pdf, MARKS_PDF, cfg())
    paginas = _textos(out)
    assert len(paginas) == 8  # 7 + 1 de índice
    assert "CONTENIDO-0" in paginas[0] and "ÍNDICE" in paginas[1]
    ent = _entradas(paginas[1])
    # Cada entrada apunta a la página (1-based) cuyo contenido es el de su primera página.
    assert "CONTENIDO-1" in paginas[ent["Datos generales"] - 1]
    assert "CONTENIDO-2" in paginas[ent["Tramo 1 › E1"] - 1]
    assert "CONTENIDO-4" in paginas[ent["Tramo 1 › E2"] - 1]
    assert "CONTENIDO-5" in paginas[ent["Recuento general"] - 1]
    assert [e["key"] for e in estructura] == ["gd", "node:1", "node:2", "recuento"]


def test_pdf_reordena_secciones_y_el_indice_sigue_el_nuevo_orden():
    pdf = _pdf_sintetico(7)
    conf = cfg(order=["recuento", "node:2", "node:1", "gd"], titles={"node:1": "Estación Uno"}, hidden=["gd"])
    out, _ = _apply(pdf, MARKS_PDF, conf)
    paginas = _textos(out)
    ent = _entradas(paginas[1])
    assert "Datos generales" not in ent  # oculta: sin entrada en el índice…
    assert not any("CONTENIDO-1" in p for p in paginas)  # …y sus páginas no se imprimen
    # orden físico de las secciones: recuento (5,6) → E2 (4) → E1 (2,3); gd (1) oculta
    fisico = [t.split("CONTENIDO-")[1][:1] for t in paginas[2:]]
    assert fisico == ["5", "6", "4", "2", "3"]
    assert "CONTENIDO-5" in paginas[ent["Recuento general"] - 1]
    assert "CONTENIDO-4" in paginas[ent["Tramo 1 › E2"] - 1]
    assert "CONTENIDO-2" in paginas[ent["Estación Uno"] - 1]  # título renombrado
    assert list(ent) == ["Recuento general", "Tramo 1 › E2", "Estación Uno"]


def test_pdf_indice_deshabilitado_no_agrega_paginas_pero_si_reordena():
    pdf = _pdf_sintetico(7)
    out, _ = _apply(pdf, MARKS_PDF, cfg(enabled=False, order=["recuento", "node:2", "node:1", "gd"]))
    paginas = _textos(out)
    assert len(paginas) == 7
    assert [t.split("CONTENIDO-")[1][:1] for t in paginas[1:]] == ["5", "6", "4", "2", "3", "1"]


def test_pdf_indice_largo_se_reparte_en_varias_paginas():
    n_sec = 60
    marks = [(f"n:{i}", f"Nodo numero {i}", 1 + i) for i in range(n_sec)]
    out, _ = _apply(_pdf_sintetico(1 + n_sec), marks, cfg())
    paginas = _textos(out)
    n_idx = len(paginas) - (1 + n_sec)
    assert n_idx >= 3
    ent = {}
    for t in paginas[1:1 + n_idx]:
        ent.update(_entradas(t))
    assert len(ent) == n_sec
    # La entrada del nodo i apunta a su contenido real (CONTENIDO-(1+i)).
    for i in (0, 17, 59):
        assert f"CONTENIDO-{1 + i}" in paginas[ent[f"Nodo numero {i}"] - 1]


def test_pdf_sin_marcas_devuelve_el_pdf_intacto():
    pdf = _pdf_sintetico(3)
    out, estructura = _apply(pdf, [], cfg())
    assert out == pdf and estructura == []


def test_pdf_con_plantilla_el_indice_hereda_el_lienzo(tmp_path):
    tpl = tmp_path / "tpl.pdf"
    c = _canvas.Canvas(str(tpl), pagesize=PAGE)
    for marca in ("PORTADA", "MAPA", "LIENZO"):
        c.drawString(50, 50, f"TPL-{marca}")
        c.showPage()
    c.save()
    out, _ = _apply(_pdf_sintetico(7), MARKS_PDF, cfg(), tpl=str(tpl))
    paginas = _textos(out)
    assert "TPL-LIENZO" in paginas[1] and "ÍNDICE" in paginas[1]


def test_pdf_outline_refleja_el_orden_final():
    out, _ = _apply(_pdf_sintetico(7), MARKS_PDF, cfg(order=["recuento", "node:2", "node:1", "gd"]))
    titulos = [o.title for o in PdfReader(io.BytesIO(out)).outline]
    assert titulos == ["ÍNDICE", "Recuento general", "Tramo 1 › E2", "Tramo 1 › E1", "Datos generales"]


# ── PPTX ───────────────────────────────────────────────────────────────────
def _prs(n_slides, con_membrete_aparte):
    """Presentación sintética: slide 0 = portada, [slide 1 = membrete], luego
    `n_slides` de contenido con el texto 'C-i'."""
    from pptx import Presentation
    from pptx.util import Cm

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(29.7), Cm(21.0)
    blank = prs.slide_layouts[6]

    def nueva(txt):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(0, 0, Cm(5), Cm(1))
        tb.text_frame.text = txt
        return s

    nueva("PORTADA")
    if con_membrete_aparte:
        nueva("MEMBRETE")
    for i in range(n_slides):
        nueva(f"C-{i}")
    return prs


def _txt(slide):
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def _aplicar_pptx(prs, marks, conf, master_separate):
    def clone():
        s = prs.slides.add_slide(prs.slide_layouts[6])
        return s

    def fill(slide, chunk, primera):
        tb = slide.shapes.add_textbox(0, 0, 100, 100)
        tb.text_frame.text = ("INDICE\n" if primera else "") + "\n".join(f"{t}|{n}" for t, n in chunk)

    return _pptx_apply_indice(prs, marks, conf, master_separate=master_separate, clone_membrete=clone,
                              fill_index_slide=fill, first_cap=4, next_cap=6)


def _entradas_pptx(slides):
    out = {}
    for s in slides:
        for ln in _txt(s).splitlines():
            if "|" in ln:
                t, n = ln.rsplit("|", 1)
                out[t] = int(n)
    return out


@pytest.mark.parametrize("membrete", [True, False])
def test_pptx_reordena_e_indice_apunta_a_la_diapositiva_real(membrete):
    base = 2 if membrete else 1  # primera diapositiva de contenido
    prs = _prs(6, membrete)
    marks = [("gd", "Datos generales", base), ("node:1", "Nodo 1", base + 1),
             ("node:2", "Nodo 2", base + 3), ("recuento", "Recuento", base + 4)]
    conf = cfg(order=["recuento", "node:2", "node:1", "gd"], titles={"node:2": "Dos"})
    est = _aplicar_pptx(prs, marks, conf, master_separate=membrete)
    slides = list(prs.slides)
    txt = [_txt(s) for s in slides]
    assert "MEMBRETE" not in "".join(txt)  # el maestro se eliminó
    assert len(slides) == 1 + 1 + 6  # portada + 1 índice + contenido (nada se perdió)
    assert "PORTADA" in txt[0] and "INDICE" in txt[1]
    contenido = [t.split("C-")[1] for t in txt[2:]]
    # recuento = contenido 4..5 ; nodo2 = 3 ; nodo1 = 1..2 ; gd = 0 (en índices de C-i)
    assert contenido == ["4", "5", "3", "1", "2", "0"]
    ent = _entradas_pptx(slides)
    for titulo, primer in (("Recuento", "4"), ("Dos", "3"), ("Nodo 1", "1"), ("Datos generales", "0")):
        assert f"C-{primer}" in txt[ent[titulo] - 1]
    assert [e["key"] for e in est] == ["gd", "node:1", "node:2", "recuento"]


def test_pptx_plantilla_de_una_slide_no_borra_contenido():
    """Sin membrete aparte, el slide de índice 1 es contenido real (antes el
    motor lo borraba siempre)."""
    prs = _prs(3, False)
    marks = [("a", "A", 1), ("b", "B", 2)]
    _aplicar_pptx(prs, marks, cfg(enabled=False), master_separate=False)
    txt = [_txt(s) for s in prs.slides]
    assert [t.split("C-")[1] for t in txt[1:]] == ["0", "1", "2"]


def test_pptx_sin_marcas_solo_retira_el_membrete():
    prs = _prs(2, True)
    est = _aplicar_pptx(prs, [], cfg(), master_separate=True)
    txt = [_txt(s) for s in prs.slides]
    assert est == [] and len(txt) == 3 and "MEMBRETE" not in "".join(txt)


def test_pptx_seccion_oculta_no_se_imprime():
    prs = _prs(6, False)
    marks = [("gd", "Datos generales", 1), ("node:1", "Nodo 1", 2), ("node:2", "Nodo 2", 4)]
    _aplicar_pptx(prs, marks, cfg(hidden=["gd"]), master_separate=False)
    txt = [_txt(s) for s in prs.slides]
    assert not any("C-0" in t for t in txt)  # la diapositiva de "gd" (C-0) no se imprime
    assert len(txt) == 1 + 1 + 5  # portada + índice + las 5 de los nodos
    ent = _entradas_pptx(list(prs.slides))
    assert "Datos generales" not in ent and set(ent) == {"Nodo 1", "Nodo 2"}


def test_fecha_de_portada_larga():
    from server import _fecha_portada_larga
    from datetime import datetime, timezone
    d = datetime(2026, 10, 16, 18, 0, tzinfo=timezone.utc)
    assert _fecha_portada_larga(d, d) == "16 DE OCTUBRE DE 2026"
    assert _fecha_portada_larga(datetime(2026, 10, 12, 18, tzinfo=timezone.utc), d) == "DEL 12 AL 16 DE OCTUBRE DE 2026"


def test_pptx_fallo_al_dibujar_deja_el_documento_como_antes():
    prs = _prs(4, True)
    marks = [("a", "A", 2), ("b", "B", 4)]

    def boom(*_a, **_k):
        raise ValueError("fallo de dibujo")

    with pytest.raises(ValueError):
        _pptx_apply_indice(prs, marks, cfg(), master_separate=True,
                           clone_membrete=lambda: prs.slides.add_slide(prs.slide_layouts[6]),
                           fill_index_slide=boom, first_cap=4, next_cap=6)
    txt = [_txt(s) for s in prs.slides]
    assert len(txt) == 5 and "MEMBRETE" not in "".join(txt)  # sin shells huérfanos
    assert [t.split("C-")[1] for t in txt[1:]] == ["0", "1", "2", "3"]


def test_periodos_son_independientes_en_la_clave_de_config():
    f = server._indice_periodo_key
    assert f("diario", "today") == "diario"
    assert f("semanal", "today") == "semanal"
    assert f("mensual", "today") == "mensual"
    assert f(None, "week") == "semanal" and f(None, "month") == "mensual" and f(None, "yesterday") == "diario"


# ── Nombres de Datos Generales ─────────────────────────────────────────────
def test_gd_names_siempre_alineados_con_las_imagenes():
    f = server._gd_names_aligned
    assert f(None, 3) == ["", "", ""]
    assert f(["  Presupuesto   anual ", "Programa"], 3) == ["Presupuesto anual", "Programa", ""]
    assert f(["a", "b", "c"], 2) == ["a", "b"]
    assert f(["x" * 500], 1) == ["x" * 120]
    assert f(["a"], 0) == []


def test_datos_generales_y_actividades_por_area_son_secciones_propias_en_el_indice():
    # Una marca por imagen de Datos Generales y una por Área: cada una con su
    # entrada, y se pueden reordenar de forma independiente.
    marks = [("gd:0", "Presupuesto", 1), ("gd:1", "Programa", 2), ("aa:OBRA CIVIL", "Actividades por área · OBRA CIVIL", 3),
             ("aa:SEGURIDAD", "Actividades por área · SEGURIDAD", 5)]
    out, est = _apply(_pdf_sintetico(7), marks, cfg(order=["aa:SEGURIDAD", "gd:1", "gd:0", "aa:OBRA CIVIL"]))
    ent = _entradas(_textos(out)[1])
    assert list(ent) == ["Actividades por área · SEGURIDAD", "Programa", "Presupuesto", "Actividades por área · OBRA CIVIL"]
    assert [e["key"] for e in est] == ["gd:0", "gd:1", "aa:OBRA CIVIL", "aa:SEGURIDAD"]
